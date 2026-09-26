"""SOV parser v2 — header-band detection, synonym map, scale detection,
totals / placeholder handling, hidden sheet & column awareness. Every value
keeps its exact cell anchor."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from uwc.refdata import CONSTRUCTION_SYNONYMS, OCCUPANCY_SYNONYMS

METHOD = "SOV parser v2 (header-band detection · synonym map · scale & totals detection)"

SYNONYMS = {
    "loc_no": ["loc #", "loc no", "site", "location #", "loc", "location number"],
    "bldg_no": ["bldg #", "building", "bldg", "bldg no"],
    "location_name": ["location name", "site name", "site description", "description", "name"],
    "address": ["street address", "address", "location address", "street"],
    "city": ["city"],
    "state": ["state", "st"],
    "zip": ["zip", "postal code", "zip code"],
    "occupancy_raw": ["occupancy", "occupancy description", "use", "occupancy / use"],
    "construction_raw": ["construction", "const type", "iso class / construction", "construction type"],
    "year_built": ["year built", "yr blt", "year of construction", "yob"],
    "stories": ["# stories", "stories", "no. of floors", "floors"],
    "sqft": ["sq ft", "area (sf)", "gross area", "square feet", "sf"],
    "roof_year": ["roof year", "roof updated", "roof replaced", "roof yr"],
    "sprinkler": ["sprinklered", "sprinkler %", "sprinkler protection", "sprinklers"],
    "building_value": ["building value", "bldg rc", "building (rc)", "building", "bldg value"],
    "contents_value": ["contents / bpp", "bpp", "contents", "business personal property"],
    "stock_value": ["stock", "inventory", "stock / inventory"],
    "bi_value": ["business income", "bi/ee (12 mo)", "time element", "bi", "business interruption"],
    "tiv_reported": ["tiv", "total insured value", "total", "total tiv"],
    "valuation": ["valuation", "valuation basis"],
    "storage_height_ft": ["max storage ht (ft)", "storage height", "max storage height"],
}
MONEY = {"building_value", "contents_value", "stock_value", "bi_value", "tiv_reported"}
_norm = lambda s: re.sub(r"\s+", " ", str(s).strip().lower())


@dataclass
class ParsedSov:
    rows: list[dict] = field(default_factory=list)       # each: {field: (value, cell_ref)}, "_row": n, "_range": "A6:T6"
    issues: list[dict] = field(default_factory=list)
    sheet: str = ""
    header_row: int = 0
    scale: float = 1.0
    columns: dict[str, str] = field(default_factory=dict)  # field -> column letter
    method: str = METHOD


def _match_header(text: str) -> str | None:
    t = _norm(text)
    for fld, syns in SYNONYMS.items():
        if t in syns:
            return fld
    return None


def map_occupancy(raw: str | None) -> tuple[str | None, float]:
    if not raw:
        return None, 0.0
    t = _norm(raw)
    if t in OCCUPANCY_SYNONYMS:
        return OCCUPANCY_SYNONYMS[t], 0.96
    best = None
    for k in sorted(OCCUPANCY_SYNONYMS, key=len, reverse=True):
        if k in t:
            best = OCCUPANCY_SYNONYMS[k]
            # battery storage beats warehouse when both present
            if "battery" in t or "lithium" in t:
                return "li_storage", 0.9
            return best, 0.86
    return None, 0.0


def map_construction(raw: str | None) -> tuple[int | None, float]:
    if not raw:
        return None, 0.0
    t = _norm(raw)
    if t in CONSTRUCTION_SYNONYMS:
        return CONSTRUCTION_SYNONYMS[t], 0.95
    hits = [k for k in CONSTRUCTION_SYNONYMS if re.search(rf"\b{re.escape(k)}\b", t)]
    if hits:
        k = max(hits, key=len)
        return CONSTRUCTION_SYNONYMS[k], 0.84
    return None, 0.0


def map_sprinkler(v: Any) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) / 100 if v > 1 else float(v)
    t = _norm(v)
    return {"y": 1.0, "yes": 1.0, "full": 1.0, "n": 0.0, "no": 0.0, "none": 0.0, "partial": 0.5}.get(t)


def parse_sov(path: Path) -> ParsedSov:
    wb = load_workbook(path, data_only=False)
    out = ParsedSov()
    sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            out.issues.append({"code": "HIDDEN_SHEET", "label": f"Hidden sheet '{ws.title}' not used as data", "anchor": {"sheet": ws.title, "cell": "A1"}, "severity": "LOW"})
    best = None
    for ws in sheets:
        for r in range(1, min(ws.max_row, 20) + 1):
            hits = sum(1 for c in ws[r] if c.value and _match_header(c.value))
            if hits >= 6 and (best is None or hits > best[2]):
                best = (ws, r, hits)
    if not best:
        out.issues.append({"code": "NO_HEADER", "label": "No header band found", "anchor": None, "severity": "HIGH"})
        return out
    ws, hr, _ = best
    out.sheet, out.header_row = ws.title, hr
    # scale detection in title band
    for r in range(1, hr):
        for c in ws[r]:
            if c.value and re.search(r"\$?0{3}s|in thousands|\(000s\)", str(c.value), re.I):
                out.scale = 1000.0
                out.issues.append({"code": "SCALE_000S", "label": "Values stated in $000s — scaled ×1,000", "anchor": {"sheet": ws.title, "cell": c.coordinate}, "severity": "MEDIUM"})
    cols: dict[str, int] = {}
    for c in ws[hr]:
        if not c.value:
            continue
        fld = _match_header(c.value)
        letter = get_column_letter(c.column)
        if ws.column_dimensions[letter].hidden:
            out.issues.append({"code": "HIDDEN_COLUMN", "label": f"Hidden column {letter} ('{c.value}') ignored", "anchor": {"sheet": ws.title, "cell": c.coordinate}, "severity": "LOW"})
            continue
        if fld and fld not in cols:
            cols[fld] = c.column
    out.columns = {k: get_column_letter(v) for k, v in cols.items()}
    last_col = get_column_letter(max(cols.values()))
    for r in range(hr + 1, ws.max_row + 1):
        if ws.row_dimensions[r].hidden:
            out.issues.append({"code": "HIDDEN_ROW", "label": f"Hidden row {r} excluded", "anchor": {"sheet": ws.title, "cell": f"A{r}"}, "severity": "MEDIUM"})
            continue
        vals = {f: ws.cell(row=r, column=ci) for f, ci in cols.items()}
        if all(v.value in (None, "") for v in vals.values()):
            continue
        first = ws.cell(row=r, column=1).value
        is_formula_total = any(isinstance(v.value, str) and v.value.startswith("=SUM") for v in vals.values())
        if (first and "total" in _norm(first)) or is_formula_total:
            out.issues.append({"code": "TOTALS_ROW", "label": f"Totals row {r} dropped (arithmetic of rows above)", "anchor": {"sheet": ws.title, "cell": f"A{r}", "range": f"A{r}:{last_col}{r}"}, "severity": "LOW"})
            continue
        row: dict[str, Any] = {"_row": r, "_range": f"A{r}:{last_col}{r}", "_sheet": ws.title}
        for f, cell in vals.items():
            v = cell.value
            if isinstance(v, str) and _norm(v) in ("tbd", "n/a", "na", "-", "unknown", "?"):
                out.issues.append({"code": "PLACEHOLDER", "label": f"Placeholder '{v}' in {cell.coordinate} treated as missing", "anchor": {"sheet": ws.title, "cell": cell.coordinate}, "severity": "LOW"})
                v = None
            if f in ("year_built", "roof_year") and isinstance(v, (int, float)) and (v < 1850 or v > 2027):
                out.issues.append({"code": "PLACEHOLDER", "label": f"Implausible year {v} in {cell.coordinate} treated as missing", "anchor": {"sheet": ws.title, "cell": cell.coordinate}, "severity": "LOW"})
                v = None
            if f in MONEY and isinstance(v, (int, float)):
                v = float(v) * out.scale
            row[f] = (v, cell.coordinate)
        # TIV arithmetic check
        comps = [row.get(k, (None,))[0] for k in ("building_value", "contents_value", "stock_value", "bi_value")]
        rep = row.get("tiv_reported", (None,))[0]
        if rep and all(isinstance(x, (int, float)) or x is None for x in comps):
            calc = sum(x or 0 for x in comps)
            if calc and abs(rep - calc) / calc > 0.005:
                out.issues.append({"code": "TIV_ARITH", "label": f"Row {r}: reported TIV {rep:,.0f} ≠ sum of components {calc:,.0f}", "anchor": {"sheet": ws.title, "cell": row['tiv_reported'][1]}, "severity": "MEDIUM"})
        out.rows.append(row)
    return out
