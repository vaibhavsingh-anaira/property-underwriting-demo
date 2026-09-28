"""HTTP API — conforms to web/src/api/types.ts."""
from __future__ import annotations

import email
import json
import threading
import time
from datetime import date, timedelta
from email import policy as email_policy
from pathlib import Path
from typing import Any

import yaml
from fastapi import Body, FastAPI, File, Header, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse, PlainTextResponse
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from uwc.config import CARRIER, DEMO_END, DEMO_START, WORLD_DIR, resolve_doc_path
from uwc import fulfilment as F
from uwc import playbooks as PB
from uwc import products as PR
from uwc.engine import rules as RE
from uwc.engine.evaluate import CONTRACT_GROUP, CONTRACT_LABEL, contract_views, evaluate_account, fmt_contract
from uwc.engine.views import resolved_field_payload
from uwc.ledger.fields import FIELDS, resolve, policy_of
from uwc.refdata import CONSTRUCTION, OCCUPANCY, USERS, USER_BY_ID, ZONES, PERIL_LABEL
from uwc.runtime import Runtime

app = FastAPI(title="Anaira Underwriting Control API", version="1.0")
RT = Runtime()
LOCK = threading.RLock()
RT.boot()

CT = {"pdf": "application/pdf", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "eml": "message/rfc822",
      "glb": "model/gltf-binary", "png": "image/png", "csv": "text/csv; charset=utf-8", "json": "application/json", "yaml": "text/yaml; charset=utf-8",
      "geojson": "application/geo+json", "txt": "text/plain; charset=utf-8"}


def uid_of(x_user_id: str | None) -> str:
    return x_user_id if x_user_id in USER_BY_ID else "u_maya"


def doc_path(doc: dict) -> Path:
    return resolve_doc_path(doc)


def days(a: str, b: str) -> int:
    return (date.fromisoformat(a[:10]) - date.fromisoformat(b[:10])).days


# ============================================================================ serializers
def doc_meta(doc: dict) -> dict:
    ex = None
    n = len(RT.store.by_doc.get(doc["doc_id"], []))
    if doc["format"] in ("pdf", "xlsx", "eml") or n:
        confs = [RT.store.obs[o].confidence for o in RT.store.by_doc.get(doc["doc_id"], [])]
        ex = {"status": "EXTRACTED" if n else "NOT_APPLICABLE", "fields": n, "avg_confidence": round(sum(confs) / len(confs), 3) if confs else None}
    return {"doc_id": doc["doc_id"], "account_id": doc.get("account_id"), "account_name": doc.get("account_name"), "doc_type": doc["doc_type"],
            "title": doc["title"], "filename": doc["filename"], "format": doc["format"], "received_at": doc["received_at"],
            "source_channel": doc["source_channel"], "size_bytes": doc.get("size_bytes", 0), "term": doc.get("term"), "is_sample": True, "extraction": ex}


def finding_payload(f: dict) -> dict:
    keys = ("finding_id", "account_id", "subject_type", "subject_id", "subject_label", "rule_id", "rule_version", "family", "title", "description",
            "severity", "outcome", "observed", "expected", "impact_usd", "impact_method", "confidence", "materiality_score", "material",
            "evidence_obs_ids", "conflicting_obs_ids", "status", "disposition", "pass", "created_at", "critique", "source")
    out = {k: f.get(k) for k in keys}
    if out["subject_type"] not in ("account", "location", "policy", "contract", "claim", "recommendation", "referral"):
        out["subject_type"] = "account"
    return out


def open_material(ren) -> list[dict]:
    return [f for f in ren.findings.values() if f["material"] and f["status"] in ("OPEN", "ACCEPTED", "DEFERRED")]


def materiality(fs: list[dict]) -> str:
    if any(f["severity"] == "CRITICAL" for f in fs):
        return "CRITICAL"
    if any(f["severity"] == "HIGH" for f in fs):
        return "HIGH"
    return "MEDIUM" if fs else "LOW"


def queue_item(acct: str) -> dict:
    reg = RT.systems["accounts"][acct]
    pas = RT.pas[acct]
    ren = RT.ren[acct]
    mats = open_material(ren)
    r = ren.rarc or {}
    locs = RT.store.account_locations(acct)
    tiv_exp = r.get("tiv_expiring") or sum((RT.systems["pas"][acct]["issued"]["locations"][i].get("building", 0) for i in range(0)), 0)
    notice = ren.notice or {}
    exp = pas.get("term_end") or RT.systems["pas"][acct]["term_end"]
    return {"account_id": acct, "name": reg["name"], "scenario": reg.get("scenario"), "segment": reg["segment"],
            "occupancy_family": OCCUPANCY.get(reg["occupancy_family"], ("—",))[0], "state": reg["hq_state"], "broker": reg["broker"],
            "underwriter": USER_BY_ID[reg["underwriter_id"]]["name"], "underwriter_id": reg["underwriter_id"], "expiry": exp,
            "days_to_expiry": days(exp, RT.clock), "notice_deadline": notice.get("latest_notice_date"), "days_to_notice": notice.get("days_remaining"),
            "status": ren.status, "pass": ren.pass_no, "recommended_actions": ren.recommended_actions if ren.pass_no else [],
            "impact_usd": round(sum(0 if f["control"] else f["impact_usd"] for f in mats)), "materiality": materiality(mats),
            "integrity_score": ren.integrity if ren.pass_no else 100, "finding_count": len(mats),
            "top_findings": [{"finding_id": f["finding_id"], "title": f["title"], "severity": f["severity"], "impact_usd": f["impact_usd"]}
                             for f in sorted(mats, key=lambda f: (-{"CRITICAL": 4, "HIGH": 3, "MEDIUM": 2, "LOW": 1}[f["severity"]], -f["impact_usd"]))[:3]],
            "tiv_expiring": r.get("tiv_expiring", 0), "tiv_current": r.get("tiv_renewal", 0), "premium_expiring": pas.get("premium") or RT.systems["pas"][acct]["premium"],
            "rarc": r.get("rarc") if ren.pass_no else None, "adequacy": r.get("adequacy") if ren.pass_no else None,
            "data_gaps": sum(1 for f in ren.findings.values() if f["outcome"] == "DATA_REQUEST" and f["status"] == "OPEN"),
            "locations": sum(1 for e in locs.values() if e.in_current or (e.in_prior and not RT.renewal_sov.get(acct))),
            "families": sorted({f["family"] for f in mats})}


def delta_cards(acct: str) -> list[dict]:
    ren = RT.ren[acct]
    r = ren.rarc or {}
    a = ren.acct_ctx or {}
    fs = [f for f in ren.findings.values() if f["status"] != "RESOLVED"]
    by = lambda *fam: [f for f in fs if f["family"] in fam]
    def status(lst):
        return "ACTION" if any(f["material"] for f in lst) else "WATCH" if lst else "OK"
    def fids(lst):
        return [f["finding_id"] for f in lst]
    locs = RT.store.account_locations(acct)
    new = [e for e in locs.values() if e.match_status == "NEW" and e.in_current]
    dele = [e for e in locs.values() if e.match_status in ("DELETED", "MERGED") and not e.in_current]
    occ = by("classification")
    exp_f = by("exposure", "valuation", "classification", "sov_integrity", "data_completeness", "data_integrity")
    tiv0, tiv1 = r.get("tiv_expiring", 0), r.get("tiv_renewal", 0)
    cards = [{
        "key": "exposure", "title": "Exposure", "status": status(exp_f),
        "headline": f"TIV {money(tiv0)} → {money(tiv1)} ({pct(tiv1 / tiv0 - 1 if tiv0 else 0, sign=True)})" if RT.renewal_sov.get(acct) else "No renewal SOV yet — expiring exposure carried forward",
        "metrics": [
            {"label": "Total insured value", "before": money(tiv0), "after": money(tiv1), "display": pct(tiv1 / tiv0 - 1 if tiv0 else 0, sign=True),
             "flag": "warn" if tiv0 and abs(tiv1 / tiv0 - 1) > 0.1 else "info"},
            {"label": "Locations", "before": str(sum(1 for e in locs.values() if e.in_prior)), "after": str(sum(1 for e in locs.values() if e.in_current) or sum(1 for e in locs.values() if e.in_prior)),
             "display": f"+{len(new)} new · {len(dele)} removed" if (new or dele) else "unchanged", "flag": "warn" if new else "ok",
             "finding_ids": fids(by("exposure"))},
            {"label": "Occupancy", "display": "changed" if occ else "unchanged", "flag": "breach" if occ else "ok", "finding_ids": fids(occ),
             "note": occ[0]["observed"] if occ else None},
            {"label": "Valuation", "display": f"{len(by('valuation'))} issue(s)" if by("valuation") else "adequate", "flag": "breach" if any(f["material"] for f in by("valuation")) else "ok",
             "finding_ids": fids(by("valuation"))},
        ]}]
    pr = by("pricing")
    if r:
        cards.append({"key": "pricing", "title": "Pricing", "status": status(pr),
                      "headline": f"Headline {pct(r['headline_change'], sign=True)} · RARC {pct(r['rarc'], sign=True)} · adequacy {pct(r['adequacy'])}",
                      "metrics": [
                          {"label": "Expiring premium", "display": money(r["expiring_premium"], full=True), "flag": "info"},
                          {"label": "Technical (renewal)", "before": money(r["tp_e0_t0"]), "after": money(r["tp_e1_t1"]), "display": money(r["tp_e1_t1"], full=True), "flag": "info"},
                          {"label": "Working proposal", "display": money(r["proposed_premium"], full=True), "note": (ren.working or {}).get("source"), "flag": "info"},
                          {"label": "RARC (like-for-like)", "display": pct(r["rarc"], sign=True), "flag": "breach" if r["rarc"] < -0.05 else "ok", "finding_ids": fids(pr)},
                          {"label": "Adequacy", "display": pct(r["adequacy"]), "flag": "breach" if (r["adequacy"] or 1) < r.get("adequacy_floor", 0.95) else "ok"},
                      ]})
    tf = by("cat_terms", "contract_integrity", "coverage", "protective_safeguards", "authority")
    cards.append({"key": "terms", "title": "Terms & contract", "status": status(tf),
                  "headline": f"{len(tf)} term / contract issue(s)" if tf else "Contract consistent; terms within guidelines",
                  "metrics": [{"label": f["title"], "display": f["observed"][:90], "flag": "breach" if f["material"] else "warn", "finding_ids": [f["finding_id"]]} for f in tf[:5]]
                  or [{"label": "Quote → binder → policy", "display": "consistent", "flag": "ok"}]})
    rq = by("claims", "engineering", "fire", "security", "vacancy", "construction", "water", "cat_integrity")
    claims = RT.claims.get(acct, [])
    since = [c for c in claims if c["dol"] >= RT.pas[acct].get("term_start", "9999")]
    recs = RT.recs.get(acct, {}).values()
    overdue = [x for x in recs if x["status"] in ("OPEN", "IN_PROGRESS") and x["due"] < RT.clock]
    cards.append({"key": "risk_quality", "title": "Risk quality", "status": status(rq),
                  "headline": f"{len(since)} loss(es) since bind · {len(overdue)} engineering action(s) overdue",
                  "metrics": [{"label": "Losses since bind", "display": f"{len(since)} · {money(sum(c['paid'] + c['reserve'] for c in since))}", "flag": "warn" if since else "ok"},
                              {"label": "Engineering overdue", "display": str(len(overdue)), "flag": "breach" if overdue else "ok", "finding_ids": fids(by("engineering"))}]
                  + [{"label": f["title"], "display": f["observed"][:90], "flag": "breach" if f["material"] else "warn", "finding_ids": [f["finding_id"]]} for f in rq if f["family"] not in ("engineering",)][:4]})
    ap = by("appetite", "accumulation")
    cards.append({"key": "appetite_portfolio", "title": "Appetite & portfolio", "status": status(ap),
                  "headline": ("; ".join(f["title"] for f in ap)) or "In appetite; no concentration breach",
                  "metrics": [{"label": "Class appetite", "display": "declined" if by("appetite") else "within appetite", "flag": "breach" if by("appetite") else "ok", "finding_ids": fids(by("appetite"))},
                              {"label": "Zone utilisation (post-renewal)", "display": f"{a.get('zone_name') or '—'} {pct(a.get('zone_util_post') or 0)}",
                               "flag": "breach" if (a.get("zone_util_post") or 0) > 0.9 else "ok", "finding_ids": fids(by("accumulation"))}]})
    reg = RT.systems["accounts"][acct]
    lr = RT.systems["pas"][acct]
    cards.append({"key": "retention", "title": "Retention & commercial", "status": "OK",
                  "headline": f"Tenure {reg['tenure_years']} yrs · broker {reg['broker']}",
                  "metrics": [{"label": "Tenure", "display": f"{reg['tenure_years']} years", "flag": "info"},
                              {"label": "Brokerage", "before": pct(RT.pas[acct].get("brokerage_prior", 0.15)), "after": pct(RT.pas[acct].get("brokerage_proposed", 0.15)),
                               "display": pct(r.get("net", {}).get("rarc_net", 0), sign=True) + " net RARC", "flag": "info"},
                              {"label": "Retention sensitivity", "display": "elevated at full technical increase" if r and r["tp_e1_t1"] > r["expiring_premium"] * 1.15 else "normal", "flag": "info"}]})
    return cards


def money(v: float | None, full: bool = False) -> str:
    if v is None:
        return "—"
    if full or abs(v) < 1e6:
        return f"${v:,.0f}"
    return f"${v / 1e6:,.1f}M"


def pct(v: float | None, sign: bool = False) -> str:
    if v is None:
        return "—"
    s = f"{v * 100:+.1f}%" if sign else f"{v * 100:.1f}%"
    return s.replace("-", "−")


def location_rows(acct: str) -> list[dict]:
    ren = RT.ren[acct]
    out = []
    fs = [f for f in ren.findings.values() if f["subject_type"] == "location" and f["status"] != "RESOLVED"]
    for uid, e in RT.store.account_locations(acct).items():
        s = ren.loc_ctx.get(uid)
        if s is None:
            from uwc.engine.views import location_state
            s = location_state(RT, acct, uid)
        prior_tiv = None
        pv = [o for o in RT.store.field_obs(uid, "tiv_reported") if o.term == "prior"]
        if e.in_prior:
            comps = []
            for f_ in ("building_value", "contents_value", "stock_value", "bi_value"):
                ob = [o for o in RT.store.field_obs(uid, f_) if o.term == "prior"]
                comps.append(ob[-1].value if ob else 0)
            prior_tiv = sum(comps) or (pv[-1].value if pv else None)
        cur_tiv = s.get("tiv") if (e.in_current or not RT.renewal_sov.get(acct)) else None
        flags = [{"code": f["rule_id"], "label": (f"Roof year: {f['observed'].replace('SOV roof year ', 'SOV ')}" if f["rule_id"] == "ROOF.YEAR_CONFLICT" else f["title"]),
                  "severity": f["severity"], "finding_id": f["finding_id"]} for f in fs if f["subject_id"] == uid]
        occ_prior = s.get("occupancy_prior")
        out.append({
            "location_uid": uid, "loc_no_prior": e.loc_no_prior, "loc_no_current": e.loc_no_current, "name": e.label, "address": e.address,
            "city": e.city, "state": e.state, "lat": e.lat or 0, "lon": e.lon or 0,
            "match_status": e.match_status if RT.renewal_sov.get(acct) or e.match_status != "MATCHED" else "MATCHED",
            "match_method": e.match_method, "match_score": e.match_score, "match_confirmed_by": e.match_confirmed_by,
            "tiv_prior": prior_tiv, "tiv_current": cur_tiv, "tiv_change_pct": (cur_tiv / prior_tiv - 1) if (prior_tiv and cur_tiv) else None,
            "occupancy": OCCUPANCY.get(s.get("occupancy_class"), (s.get("occupancy_raw") or "—",))[0],
            "occupancy_prior": OCCUPANCY.get(occ_prior, ("—",))[0] if occ_prior and occ_prior != s.get("occupancy_class") else None,
            "construction": CONSTRUCTION.get(s.get("construction_class") or 0, (s.get("construction_raw") or "—",))[0],
            "year_built": s.get("year_built"), "stories": s.get("stories"), "sqft": s.get("sqft"), "roof_year": s.get("roof_year"),
            "sprinkler": f"{(s.get('sprinkler_pct') or 0) * 100:.0f}%" if s.get("sprinkler_pct") is not None else "—",
            "valuation_ratio": s.get("valuation_ratio"), "wind_tier": s.get("wind_tier"), "cat_zone": ZONES.get(s.get("cat_zone"), (None,))[0] if s.get("cat_zone") else None,
            "aal": s.get("aal"), "buildings": 1, "flags": flags, "model_doc_id": e.buildings_model_doc, "imagery_doc_ids": list(e.imagery_docs),
        })
    order = {"NEW": 0, "AMBIGUOUS": 1, "MERGED": 2, "DELETED": 3, "MATCHED": 4}
    out.sort(key=lambda x: (order.get(x["match_status"], 5), -(len(x["flags"])), x["loc_no_current"] or x["loc_no_prior"] or ""))
    return out


# ============================================================================ routes: people & demo
@app.get("/api/users")
def users():
    return USERS


def demo_state() -> dict:
    nxt = []
    for e in RT.pending_events():
        if e["type"].startswith("sim.") or e["type"] in ("cat.run_published", "vendor.model3d"):
            continue
        nxt.append({"date": e["date"], "title": e["title"], "account_name": e.get("account_name")})
        if len(nxt) >= 12:
            break
    return {"clock": RT.clock, "start": DEMO_START.isoformat(), "end": DEMO_END.isoformat(), "events_applied": RT.applied,
            "events_total": len(RT.events), "next_events": nxt, "injections": RT.injections}


@app.get("/api/demo/state")
def get_demo():
    with LOCK:
        return demo_state()


@app.post("/api/demo/reset")
def post_reset():
    with LOCK:
        RT.reset()
        return demo_state()


@app.post("/api/demo/advance")
def post_advance(b: dict = Body(...)):
    with LOCK:
        if b.get("to"):
            target = b["to"]
        else:
            target = (date.fromisoformat(RT.clock) + timedelta(days=int(b.get("days", 1)))).isoformat()
        if target <= RT.clock:
            return {"clock": RT.clock, "applied": [], "new_findings": 0}
        applied = RT.advance_to(target)
        tl = [{"event_id": e.get("event_id", f"dyn_{i}"), "date": e["date"], "kind": "system", "stage": "", "title": e["title"], "detail": e.get("account_name") or "",
               "actor": "Timeline", "doc_id": None, "finding_ids": []} for i, e in enumerate(applied) if not e["type"].startswith("sim.")]
        return {"clock": RT.clock, "applied": tl, "new_findings": RT.new_findings_counter}


@app.post("/api/demo/injections")
def post_injections(b: dict = Body(...)):
    with LOCK:
        RT.injections.update({k: bool(v) for k, v in b.items()})
        return demo_state()


# ============================================================================ book
@app.get("/api/book")
def book():
    with LOCK:
        accts = list(RT.systems["accounts"])
        items = [queue_item(a) for a in accts]
        ident = val = appr = corr = 0.0
        for a in accts:
            ren = RT.ren[a]
            for f in ren.findings.values():
                if not f["material"] or f["control"]:
                    continue
                ident += f["impact_usd"]
                if f["status"] in ("ACCEPTED",) or (f["status"] == "RESOLVED" and f.get("disposition", {}) and (f.get("disposition") or {}).get("decision") == "ACCEPT"):
                    val += f["impact_usd"]
                    if any(q["status"] in ("APPROVED", "SENT", "ACCEPTED", "BOUND") for q in ren.quotes):
                        appr += f["impact_usd"]
                    if ren.r_bound and f["status"] == "RESOLVED":
                        corr += f["impact_usd"]
        quoted = [(a, RT.ren[a]) for a in accts if RT.ren[a].quotes]
        wsum = sum(RT.pas[a]["premium"] for a, _ in quoted) or 1
        rarc_c = sum(r.quotes[-1]["rarc"] * RT.pas[a]["premium"] for a, r in quoted) / wsum if quoted else 0
        rarc_r = sum(r.quotes[-1]["uw_reported_rarc"] * RT.pas[a]["premium"] for a, r in quoted) / wsum if quoted else 0
        hist = {}
        for it in items:
            if it["adequacy"] is None:
                continue
            b = min(1.2, max(0.7, round(it["adequacy"] * 20) / 20))
            k = f"{int(b * 100)}%"
            h = hist.setdefault(k, {"bucket": k, "count": 0, "premium": 0, "_b": b})
            h["count"] += 1
            h["premium"] += it["premium_expiring"]
        fam: dict[str, dict] = {}
        act: dict[str, dict] = {}
        for a in accts:
            for f in open_material(RT.ren[a]):
                x = fam.setdefault(f["family"], {"family": f["family"], "label": f["family"].replace("_", " ").capitalize(), "count": 0, "impact_usd": 0})
                x["count"] += 1
                x["impact_usd"] += 0 if f["control"] else f["impact_usd"]
            for ac in RT.ren[a].actions:
                y = act.setdefault(ac["type"], {"action": ac["type"], "count": 0, "impact_usd": 0})
                y["count"] += 1
                y["impact_usd"] += ac["impact_usd"]
        uw: dict[str, dict] = {}
        for it in items:
            u = uw.setdefault(it["underwriter_id"], {"user_id": it["underwriter_id"], "name": it["underwriter"], "accounts": 0, "findings": 0, "impact_usd": 0,
                                                     "pricing_deviations": 0, "_r": [], "avg_rarc": 0})
            u["accounts"] += 1
            u["findings"] += it["finding_count"]
            u["impact_usd"] += it["impact_usd"]
            if it["rarc"] is not None:
                u["_r"].append(it["rarc"])
                u["pricing_deviations"] += 1 if it["rarc"] < -0.05 else 0
        for u in uw.values():
            u["avg_rarc"] = sum(u["_r"]) / len(u["_r"]) if u["_r"] else 0
            del u["_r"]
        br: dict[str, dict] = {}
        for it in items:
            b_ = br.setdefault(it["broker"], {"broker": it["broker"], "accounts": 0, "premium": 0, "impact_usd": 0, "_r": []})
            b_["accounts"] += 1
            b_["premium"] += it["premium_expiring"]
            b_["impact_usd"] += it["impact_usd"]
            if it["rarc"] is not None:
                b_["_r"].append(it["rarc"])
        for b_ in br.values():
            b_["avg_rarc"] = sum(b_["_r"]) / len(b_["_r"]) if b_["_r"] else 0
            del b_["_r"]
        months: dict[str, dict] = {}
        for it in items:
            m = it["expiry"][:7]
            x = months.setdefault(m, {"month": m, "renewals": 0, "premium": 0, "fast_track": 0, "action": 0})
            x["renewals"] += 1
            x["premium"] += it["premium_expiring"]
            if it["status"] == "FAST_TRACK":
                x["fast_track"] += 1
            elif it["pass"]:
                x["action"] += 1
        return {
            "as_of": RT.clock, "carrier": CARRIER, "renewals": len(items), "locations": sum(i["locations"] for i in items),
            "tiv": sum(i["tiv_current"] or i["tiv_expiring"] for i in items), "premium_expiring": sum(i["premium_expiring"] for i in items),
            "fast_track": sum(1 for i in items if i["status"] == "FAST_TRACK" or (RT.ren[i["account_id"]].fast_track_confirmed)),
            "material_action": sum(1 for i in items if i["finding_count"] > 0), "not_started": sum(1 for i in items if i["pass"] == 0),
            "pipeline": {"identified": round(ident), "validated": round(val), "approved": round(appr), "corrected": round(corr)},
            "rarc_reported": rarc_r, "rarc_computed": rarc_c,
            "adequacy_hist": [{k: v for k, v in h.items() if k != "_b"} for h in sorted(hist.values(), key=lambda h: h["_b"])],
            "by_family": sorted(fam.values(), key=lambda x: -x["impact_usd"]), "by_action": sorted(act.values(), key=lambda x: -x["count"]),
            "by_underwriter": sorted(uw.values(), key=lambda x: -x["impact_usd"]), "by_broker": sorted(br.values(), key=lambda x: -x["premium"]),
            "by_month": sorted(months.values(), key=lambda x: x["month"]), "accumulation": accumulation(),
        }


def accumulation() -> list[dict]:
    out = []
    for z, (name, peril, lat, lon, rad, thr, bg) in ZONES.items():
        cur = RT.zone_total(z, "prior")
        post = RT.zone_total(z, "current")
        accts = sorted([{"account_id": a, "name": RT.systems["accounts"][a]["name"], "contribution": c.get("current", {}).get(z, 0)}
                        for a, c in RT.zone_contrib.items() if c.get("current", {}).get(z)], key=lambda x: -x["contribution"])[:8]
        out.append({"zone_id": z, "name": name, "peril": peril, "lat": lat, "lon": lon, "radius_km": rad, "threshold": thr, "current": cur,
                    "post_renewal": post, "utilization": post / thr, "accounts": accts})
    return sorted(out, key=lambda x: -x["utilization"])


@app.get("/api/portfolio/accumulation")
def get_accumulation():
    with LOCK:
        return accumulation()


@app.get("/api/renewals")
def renewals(status: str | None = None, uw: str | None = None, q: str | None = None):
    with LOCK:
        items = [queue_item(a) for a in RT.systems["accounts"]]
        if status:
            items = [i for i in items if i["status"] == status]
        if uw:
            items = [i for i in items if i["underwriter_id"] == uw]
        if q:
            ql = q.lower()
            items = [i for i in items if ql in i["name"].lower() or ql in i["broker"].lower()]
        rank = {"CRITICAL": 4, "HIGH": 3, "MEDIUM": 2, "LOW": 1}
        items.sort(key=lambda i: (-(i["impact_usd"] * rank[i["materiality"]]) / max(5, (i["days_to_notice"] if i["days_to_notice"] is not None else i["days_to_expiry"])), i["days_to_expiry"]))
        return items


# ============================================================================ account
@app.get("/api/accounts/{acct}")
def account(acct: str):
    with LOCK:
        if acct not in RT.systems["accounts"]:
            raise HTTPException(404, "account not found")
        reg = RT.systems["accounts"][acct]
        pas = RT.pas[acct] if RT.pas[acct].get("term_start") else {**RT.systems["pas"][acct]}
        ren = RT.ren[acct]
        cv = contract_views(RT, acct)
        pol = cv.get("endorsed") or cv.get("policy") or {}
        layer = "Primary" if not pas.get("layer_attach") else f"{money(pas['layer_limit'])} xs {money(pas['layer_attach'])}"
        if pas.get("carrier_share", 1) < 1:
            layer += f" · {pas['carrier_share'] * 100:.0f}% share"
        limit_basis = pol.get("limit_basis") or "blanket"
        missing = [str(o.value) for o in RT.store.field_obs(acct, "missing_document")]
        docs = [doc_meta(d) for d in RT.store.docs.values() if d.get("account_id") == acct]
        docs.sort(key=lambda d: d["received_at"], reverse=True)
        mats = open_material(ren)
        conf = "HIGH" if all((f.get("critique") or {}).get("verdict") == "UPHELD" for f in mats) else "MEDIUM"
        return {
            "account_id": acct, "name": reg["name"], "scenario": reg.get("scenario"), "scenario_title": reg.get("scenario_title"),
            "segment": reg["segment"], "occupancy_family": OCCUPANCY.get(reg["occupancy_family"], ("—",))[0], "broker": reg["broker"],
            "broker_contact": reg["broker_contact"], "underwriter": USER_BY_ID[reg["underwriter_id"]]["name"], "underwriter_id": reg["underwriter_id"],
            "state": reg["hq_state"], "tenure_years": reg["tenure_years"],
            "policy": {"policy_id": pas["policy_no"], "policy_no": pas["policy_no"], "term_start": pas["term_start"], "term_end": pas["term_end"],
                       "carrier_share": pas.get("carrier_share", 1.0), "layer": layer, "limit": pol.get("limit") or 0, "limit_basis": limit_basis,
                       "admitted": reg["admitted"], "premium": pas["premium"],
                       "forms": [{"form_no": f, "title": __import__("uwc.refdata", fromlist=["FORMS"]).FORMS.get(f, "")} for f in (pol.get("forms") or [])]},
            "renewal": {"expiry": pas["term_end"], "days_to_expiry": days(pas["term_end"], RT.clock), "pass": ren.pass_no, "status": ren.status,
                        "integrity_score": ren.integrity if ren.pass_no else 100, "recommended_actions": ren.recommended_actions if ren.pass_no else [],
                        "confidence": conf, "owner": USER_BY_ID[reg["underwriter_id"]]["name"],
                        "due": (date.fromisoformat(pas["term_end"]) - timedelta(days=60)).isoformat(),
                        "notice": ren.notice or {"required": False, "state": reg["hq_state"], "admitted": reg["admitted"], "days_required": None,
                                                 "latest_notice_date": None, "days_remaining": None, "rule_ref": ""},
                        "missing": missing},
            "deltas": delta_cards(acct) if ren.pass_no else [],
            "findings": [finding_payload(f) for f in ren.findings.values()],
            "actions": [{k: v for k, v in a.items()} for a in ren.actions],
            "narrative": ren.narrative or {"text": "Pass 1 has not run yet (starts at T-150).", "generated_by": "template", "model": None, "critique": None},
            "locations": location_rows(acct), "timeline": list(reversed(ren.timeline)), "documents": docs,
            "site_model_doc_id": RT.site_models.get(acct),
        }


@app.get("/api/accounts/{acct}/fields")
def fields(acct: str, subject_id: str | None = None):
    with LOCK:
        sid = subject_id or acct
        return [resolved_field_payload(RT, sid, f) for f in RT.store.subject_fields(sid)]


ALIAS = {"tiv": "tiv_reported", "roof": "roof_year", "occ": "occupancy_class", "const": "construction_class"}


@app.get("/api/observations/{obs_id}")
def observation(obs_id: str):
    with LOCK:
        if obs_id.startswith("ob-"):
            uid, _, key = obs_id[3:].rpartition("-")
            fld = ALIAS.get(key, key)
            obs = RT.store.field_obs(uid, fld)
            if fld == "tiv_reported":
                obs = [o for o in obs if o.term == "current"] or obs
            w, _ = resolve(obs, policy_of(fld)) if obs else (None, False)
            if not w:
                raise HTTPException(404, "no evidence for this field")
            obs_id = w.obs_id
        o = RT.store.obs.get(obs_id)
        if not o:
            raise HTTPException(404, "observation not found")
        field = resolved_field_payload(RT, o.subject_id, o.field_code)
        winner = RT.store.obs.get(field["resolved_obs_id"]) if field["resolved_obs_id"] else None
        doc = RT.store.docs.get(o.doc_id) or (RT.world_docs.get(o.anchor.get("doc_id")) if o.anchor and o.anchor.get("doc_id") else None)
        return {"observation": RT.obs_payload(o, winner), "field": field, "document": doc_meta(doc) if doc else None}


def rarc_payload(acct: str, r: dict) -> dict:
    keys = ("expiring_premium", "tp_e0_t0", "tp_e1_t0", "tp_e1_t1", "exposure_factor", "terms_factor", "expected_premium", "proposed_premium",
            "headline_change", "rarc", "adequacy", "adequacy_floor", "price_deviation", "required_authority_level", "authority_reasons", "method",
            "model_version", "model_drift", "net", "breakdown", "tiv_expiring", "tiv_renewal", "expiring_terms", "proposed_terms")
    return {"account_id": acct, **{k: r.get(k) for k in keys}}


@app.get("/api/accounts/{acct}/rarc")
def rarc(acct: str):
    with LOCK:
        r = RT.ren[acct].rarc
        if not r:
            raise HTTPException(409, "Pass 1 has not run for this account yet")
        return rarc_payload(acct, r)


@app.post("/api/accounts/{acct}/rarc/whatif")
def whatif(acct: str, b: dict = Body(...)):
    with LOCK:
        return rarc_payload(acct, RT.whatif(acct, float(b["proposed_premium"]), b.get("terms") or {}, b.get("brokerage")))


@app.get("/api/accounts/{acct}/contract")
def contract(acct: str):
    with LOCK:
        ren = RT.ren[acct]
        pas = RT.pas[acct]
        cv = contract_views(RT, acct)
        renewal = bool(ren.r_bound)
        if renewal:
            cv = {**{k: v for k, v in contract_views(RT, acct, renewal=True).items()}}
        labels = {uid: e.label for uid, e in RT.store.account_locations(acct).items()}
        rows = []
        fields = ["limit", "limit_basis", "aop_deductible", "named_storm_ded_pct", "named_storm_ded_min", "wind_hail_ded_pct", "bi_sublimit",
                  "flood_sublimit", "eq_sublimit", "premium", "forms", "safeguards"]
        fmap = {f["subject_id"].split(":", 2)[-1].replace(":r", ""): f["finding_id"] for f in ren.findings.values()
                if f["family"] == "contract_integrity" and f["status"] != "RESOLVED" and f["subject_type"] == "contract"}
        for f in fields:
            vals, anchors = {}, {}
            for col in ("quote", "binder", "policy", "endorsed"):
                v = cv.get(col)
                if v is None:
                    continue
                raw = v.get(f)
                if f == "safeguards" and raw:
                    vals[col] = ", ".join(f"{x.split('|')[1]} {labels.get(x.split('|')[0], '')}".strip() for x in raw)
                elif f == "forms" and raw:
                    vals[col] = ", ".join(raw)
                elif f == "limit_basis":
                    vals[col] = {"blanket": "Blanket", "scheduled": "Scheduled", "loss_limit": "Loss limit"}.get(raw, raw)
                else:
                    vals[col] = fmt_contract(f, raw) or "—"
                oid = v["_obs"].get(f)
                anchors[col] = RT.store.obs[oid].anchor if oid and oid in RT.store.obs else None
            present = [vals.get(c) for c in ("quote", "binder", "policy") if c in cv]
            if f == "premium" or (f in ("forms", "safeguards") and "quote" in cv and vals.get("quote") is None):
                present = [vals.get(c) for c in ("binder", "policy") if c in cv]
            result = "MATCH"
            if len(set(present)) > 1:
                result = "MISSING" if any(x is None for x in present) else "MISMATCH"
            fid = fmap.get(f"binder_policy:{f}") or fmap.get(f"quote_binder:{f}")
            rows.append({"field": f, "label": CONTRACT_LABEL.get(f, f), "group": CONTRACT_GROUP.get(f, "Limits"), "values": vals, "anchors": anchors,
                         "result": result if f != "premium" else ("MATCH" if len(set(present)) <= 1 else "MISMATCH"), "finding_id": fid})
        subj = []
        for i, sj in enumerate(pas.get("subjectivities", [])):
            fid = next((f["finding_id"] for f in ren.findings.values() if f["subject_id"] == f"{acct}:subj:{i}" and f["status"] != "RESOLVED"), None)
            subj.append({"text": sj["text"], "due": sj["due"], "status": sj["status"], "age_days": days(RT.clock, pas.get("bind_date") or RT.clock), "finding_id": fid})
        endts = [{"endt_id": e["endt_id"], "effective": e["effective"], "type": e["type"], "description": e["description"], "premium_delta": e["premium_delta"],
                  "doc_id": e.get("doc_id")} for e in pas.get("endorsements_applied", [])]
        quotes = []
        for q in pas.get("quotes", []):
            quotes.append({"quote_id": f"{acct}_prior_q{q['version']}", "version": q["version"], "term": f"{pas['term_start'][:4]}–{pas['term_end'][:4]}",
                           "created_at": q["date"], "created_by": USER_BY_ID[RT.systems["accounts"][acct]["underwriter_id"]]["name"], "premium": q["premium"],
                           "terms": {k: q["terms"].get(k) for k in ("aop_deductible", "named_storm_ded_pct", "named_storm_ded_min", "wind_hail_ded_pct", "bi_sublimit", "flood_sublimit", "eq_sublimit")},
                           "status": "BOUND" if q["status"] == "ACCEPTED" else "SUPERSEDED", "terms_hash": __import__("uwc.engine.pricing", fromlist=["x"]).terms_hash(q["premium"], q["terms"]),
                           "doc_id": q.get("doc_id"), "rarc": None, "adequacy": None})
        for q in ren.quotes:
            quotes.append({k: q.get(k) for k in ("quote_id", "version", "term", "created_at", "created_by", "premium", "terms", "status", "terms_hash", "doc_id", "rarc", "adequacy")})
        refs = [referral_payload(x) for x in ren.referrals]
        for ap in pas.get("approvals", []):
            refs.insert(0, {"referral_id": f"{acct}_prior_ref_{ap['quote_version']}", "account_id": acct, "account_name": RT.systems["accounts"][acct]["name"],
                            "quote_id": f"{acct}_prior_q{ap['quote_version']}", "terms_hash": None, "requested_by": USER_BY_ID[RT.systems["accounts"][acct]["underwriter_id"]]["name"],
                            "requested_at": ap["date"], "required_level": 3, "reasons": ["Prior-term referral"], "memo": "Prior-term approval (PAS record)",
                            "status": "APPROVED", "approver": USER_BY_ID[ap["approver"]]["name"], "decided_at": ap["date"], "conditions": None,
                            "invalidated_reason": None})
            bound = next((q for q in pas.get("quotes", []) if q["status"] == "ACCEPTED"), None)
            if bound and bound["version"] != ap["quote_version"]:
                refs[0]["status"] = "INVALIDATED"
                refs[0]["invalidated_reason"] = f"Terms changed after approval (v{ap['quote_version']} approved → v{bound['version']} bound)"
        return {"term": f"{pas['term_start'][:4]}–{pas['term_end'][:4]}" if not renewal else "renewal", "quote_version": str(next((q["version"] for q in pas.get("quotes", []) if q["status"] == "ACCEPTED"), "—")),
                "rows": rows, "endorsements": endts, "subjectivities": subj, "quotes": quotes, "referrals": refs}


def referral_payload(x: dict) -> dict:
    keys = ("referral_id", "account_id", "account_name", "quote_id", "terms_hash", "requested_by", "requested_at", "required_level", "reasons", "memo",
            "status", "approver", "decided_at", "conditions", "invalidated_reason")
    return {k: x.get(k) for k in keys}


@app.get("/api/accounts/{acct}/cat")
def cat(acct: str):
    with LOCK:
        runs = RT.cat_runs.get(acct, [])
        order = {"CURRENT": 0, "RENEWAL_PROPOSED": 1, "AS_BOUND": 2}
        out = []
        labels = {uid: e.label for uid, e in RT.store.account_locations(acct).items()}
        ren = RT.ren[acct]
        mism = []
        for uid, s in ren.loc_ctx.items():
            if s.get("cat_input_construction") and s.get("construction_class") and s["cat_input_construction"] != s["construction_class"]:
                mism.append({"location_uid": uid, "field": "Construction", "cat_input": CONSTRUCTION[s["cat_input_construction"]][0],
                             "resolved": CONSTRUCTION[s["construction_class"]][0]})
        compare = [{"label": f"{r['snapshot'].replace('_', ' ').title()} · {r['run_date']}", "aal_total": r["result"]["aal_total"],
                    "oep_100": next((x["loss"] for x in r["result"]["oep"] if x["rp"] == 100), 0), "oep_250": next((x["loss"] for x in r["result"]["oep"] if x["rp"] == 250), 0)} for r in runs]
        from uwc.mocks.cat import dq_flags
        for r in sorted(runs, key=lambda r: order.get(r["snapshot"], 9)):
            res = r["result"]
            la = res.get("location_aal", {})
            tot = sum(la.values()) or 1
            expo = {x["location_uid"]: x for x in r.get("exposure", [])}
            contrib = []
            for k, v in sorted(la.items(), key=lambda kv: -kv[1]):
                uid = k if k in labels else next((u for u, e in RT.store.account_locations(acct).items() if e.loc_no_prior == k.lstrip("L")), k)
                e = expo.get(k, {})
                contrib.append({"location_uid": uid, "label": labels.get(uid, k), "aal": v, "pct": v / tot, "tiv": sum(e.get(x, 0) for x in ("building", "contents", "stock", "bi"))})
            flags = dq_flags(r.get("exposure", [])) if r.get("exposure") else []
            for fl in flags:
                fl["label"] = labels.get(fl["location_uid"], fl["label"])
            out.append({"run_id": r["ep_doc_id"], "account_id": acct, "snapshot": r["snapshot"], "vendor": "MockCat (stand-in for Moody's RMS / Verisk)",
                        "model_version": r["model_version"], "run_date": r["run_date"], "perils": [PERIL_LABEL.get(p, p) for p in res.get("aal_by_peril", {})],
                        "basis": res.get("basis", "Carrier share, net of deductibles"), "aal_total": res["aal_total"],
                        "aal_by_peril": [{"peril": PERIL_LABEL.get(k, k), "aal": v} for k, v in res.get("aal_by_peril", {}).items()],
                        "oep": res["oep"], "aep": res["aep"], "location_contrib": contrib, "dq_flags": flags,
                        "input_mismatches": mism if r["snapshot"] == "AS_BOUND" else [], "exposure_doc_id": r["exposure_doc_id"], "elt_doc_id": r["elt_doc_id"],
                        "ep_doc_id": r["ep_doc_id"], "compare": [c for c in compare if not c["label"].startswith(r["snapshot"].replace("_", " ").title())]})
        return out


@app.get("/api/accounts/{acct}/claims-engineering")
def claims_eng(acct: str):
    with LOCK:
        labels = {uid: e.label for uid, e in RT.store.account_locations(acct).items()}
        cl = RT.claims.get(acct, [])
        claims = [{"claim_id": c["claim_id"], "location_uid": c.get("location_uid"), "location_label": labels.get(c.get("location_uid"), c.get("location_address") or "—"),
                   "date_of_loss": c["dol"], "cause": c["cause"].replace("_", " "), "cat_event": c.get("cat_event"), "status": c["status"], "paid": c["paid"],
                   "reserve": c["reserve"], "incurred": c["paid"] + c["reserve"], "description": c["description"], "linked_recommendation": c.get("linked_rec")} for c in cl]
        by: dict[str, dict] = {}
        five = [c for c in claims if c["date_of_loss"] >= f"{int(RT.clock[:4]) - 5}{RT.clock[4:]}"]
        for c in five:
            x = by.setdefault(c["cause"], {"cause": c["cause"], "count": 0, "incurred": 0})
            x["count"] += 1
            x["incurred"] += c["incurred"]
        prem = RT.pas[acct].get("premium") or 1
        recs = []
        for r in RT.recs.get(acct, {}).values():
            od = max(0, days(RT.clock, r["due"])) if r["status"] not in ("VERIFIED_CLOSED",) else 0
            recs.append({"rec_id": r["rec_id"], "location_uid": r.get("location_uid"), "location_label": labels.get(r.get("location_uid"), ""), "raised": r["raised"],
                         "category": r["category"], "description": r["description"], "severity": r["severity"], "due": r["due"], "status": r["status"],
                         "bind_condition": r["bind_condition"], "days_overdue": od if r["status"] in ("OPEN", "IN_PROGRESS") or (r["bind_condition"] and r["status"] != "VERIFIED_CLOSED") else 0,
                         "completion_evidence": r.get("completion_evidence"), "doc_id": None})
        surveys = []
        for d_ in RT.store.docs.values():
            if d_.get("account_id") == acct and d_["doc_type"] == "Engineering report":
                surveys.append({"survey_id": d_["doc_id"], "location_label": "All locations", "date": d_["title"].split("— ")[-1], "engineer": "Elena Brooks, CSP", "doc_id": d_["doc_id"]})
        return {"claims": claims, "summary": {"count_5y": len(five), "incurred_5y": sum(c["incurred"] for c in five),
                                               "loss_ratio_5y": sum(c["incurred"] for c in five) / (prem * 5), "by_cause": sorted(by.values(), key=lambda x: -x["incurred"])},
                "recommendations": recs, "surveys": surveys}


# ============================================================================ workflow actions
@app.post("/api/findings/{fid}/disposition")
def disposition(fid: str, b: dict = Body(...), x_user_id: str | None = Header(None)):
    with LOCK:
        try:
            f = RT.disposition(fid, b["decision"], b.get("reason_code", "agree"), b.get("note", ""), uid_of(x_user_id))
        except KeyError:
            raise HTTPException(404, "finding not found")
        return finding_payload(f)


def quote_payload(q: dict) -> dict:
    return {k: q.get(k) for k in ("quote_id", "version", "term", "created_at", "created_by", "premium", "terms", "status", "terms_hash", "doc_id", "rarc", "adequacy")}


@app.post("/api/accounts/{acct}/confirm")
def confirm(acct: str, b: dict = Body(...), x_user_id: str | None = Header(None)):
    with LOCK:
        ren = RT.ren[acct]
        ren.fast_track_confirmed = True
        RT.log(acct, "user", "15", "Fast-track confirmed — maintain", "No material change; renewal on expiring terms", USER_BY_ID[uid_of(x_user_id)]["name"])
        return queue_item(acct)


@app.post("/api/accounts/{acct}/quotes")
def create_quote(acct: str, b: dict = Body(...), x_user_id: str | None = Header(None)):
    with LOCK:
        try:
            return quote_payload(RT.create_quote(acct, float(b["premium"]), b.get("terms") or {}, uid_of(x_user_id)))
        except ValueError as e:
            raise HTTPException(409, str(e))


@app.post("/api/accounts/{acct}/quotes/{qid}/send")
def send_quote(acct: str, qid: str, x_user_id: str | None = Header(None)):
    with LOCK:
        try:
            return quote_payload(RT.send_quote(acct, qid, uid_of(x_user_id)))
        except PermissionError as e:
            raise HTTPException(403, str(e))


@app.post("/api/accounts/{acct}/referrals")
def create_referral(acct: str, b: dict = Body(...), x_user_id: str | None = Header(None)):
    with LOCK:
        return referral_payload(RT.create_referral(acct, b.get("quote_id"), b.get("note", ""), uid_of(x_user_id)))


@app.get("/api/referrals")
def referrals(status: str | None = None):
    with LOCK:
        out = [referral_payload(x) for r in RT.ren.values() for x in r.referrals]
        if status:
            out = [x for x in out if x["status"] == status]
        return sorted(out, key=lambda x: x["requested_at"], reverse=True)


@app.post("/api/referrals/{rid}/decision")
def decide(rid: str, b: dict = Body(...), x_user_id: str | None = Header(None)):
    with LOCK:
        try:
            return referral_payload(RT.decide_referral(rid, b["decision"], b.get("conditions"), uid_of(x_user_id)))
        except PermissionError as e:
            raise HTTPException(403, str(e))
        except KeyError:
            raise HTTPException(404, "referral not found")


@app.post("/api/accounts/{acct}/data-request")
def data_request(acct: str, b: dict = Body(...), x_user_id: str | None = Header(None)):
    with LOCK:
        return RT.data_request(acct, b.get("items", []), uid_of(x_user_id))


@app.post("/api/accounts/{acct}/bind")
def bind(acct: str, b: dict = Body(...), x_user_id: str | None = Header(None)):
    with LOCK:
        try:
            r = RT.bind(acct, b["quote_id"], uid_of(x_user_id))
        except (ValueError, PermissionError) as e:
            raise HTTPException(409, str(e))
        return {"binder_doc_id": r["binder_doc_id"], "findings": [finding_payload(f) for f in r["findings"]]}


@app.post("/api/accounts/{acct}/issue")
def issue(acct: str, x_user_id: str | None = Header(None)):
    with LOCK:
        try:
            r = RT.issue(acct, uid_of(x_user_id))
        except ValueError as e:
            raise HTTPException(409, str(e))
        return {"policy_doc_id": r["policy_doc_id"], "findings": [finding_payload(f) for f in r["findings"]]}


# ============================================================================ fulfilment (clearance, broker docs, engineering, subjectivities, billing)
def fulfilment_payload(acct: str) -> dict:
    ren = RT.ren[acct]
    pend = [{"date": e["date"], "type": e["type"], "title": e["title"]} for e in RT.dynamic if e.get("account_id") == acct and not e.get("done")]
    return {"clearance": F.clearance(RT, acct), "renewal_subjectivities": F.renewal_subjectivities(RT, acct), "received": F.received(RT, acct),
            "billing": F._st(RT)["billing"].get(acct), "pending": pend, "recommendations": list(RT.recs.get(acct, {}).values()),
            "issued": ren.r_issued, "bound": {k: v for k, v in (ren.r_bound or {}).items() if k != "terms"} or None}


@app.get("/api/accounts/{acct}/fulfilment")
def fulfilment(acct: str):
    with LOCK:
        return fulfilment_payload(acct)


@app.post("/api/accounts/{acct}/engineering/survey")
def order_survey(acct: str, b: dict = Body(default={}), x_user_id: str | None = Header(None)):
    with LOCK:
        F.order_survey(RT, acct, uid_of(x_user_id), b.get("scope") or "Renewal verification survey")
        return fulfilment_payload(acct)


@app.post("/api/accounts/{acct}/engineering/recs/{rec_id}/verify")
def verify_rec(acct: str, rec_id: str, x_user_id: str | None = Header(None)):
    with LOCK:
        F.verify_rec(RT, acct, rec_id, uid_of(x_user_id))
        evaluate_account(RT, acct)
        return fulfilment_payload(acct)


@app.post("/api/accounts/{acct}/issue/correct")
def correct_issue(acct: str, x_user_id: str | None = Header(None)):
    with LOCK:
        if not RT.ren[acct].r_issued:
            raise HTTPException(409, "Nothing issued yet")
        return {"message": F.correct_issuance(RT, acct, uid_of(x_user_id)), **fulfilment_payload(acct)}


@app.post("/api/accounts/{acct}/events")
def presenter_event(acct: str, b: dict = Body(...), x_user_id: str | None = Header(None)):
    with LOCK:
        try:
            msg = F.presenter_event(RT, acct, b["kind"], b.get("params") or {}, uid_of(x_user_id))
        except ValueError as e:
            raise HTTPException(400, str(e))
        return {"message": msg}


# ============================================================================ sandbox — your own SOV through the real parser and controls
@app.post("/api/sandbox/sov")
async def sandbox_sov(renewal: UploadFile = File(...), expiring: UploadFile | None = File(None)):
    from uwc import sandbox as SB
    for f in (renewal, expiring):
        if f and not f.filename.lower().endswith((".xlsx", ".csv")):
            raise HTTPException(400, f"{f.filename}: upload .xlsx or .csv")
    try:
        return SB.analyse((renewal.filename, await renewal.read()), (expiring.filename, await expiring.read()) if expiring else None)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(422, f"Could not read the schedule: {e}")


@app.get("/api/sandbox/samples/{name}")
def sandbox_sample(name: str):
    files = {"expiring": "d_s2_sov_2025.xlsx", "renewal": "d_s2_sov_2026.xlsx", "messy": "d_s11_sov_2026.xlsx"}
    if name not in files:
        raise HTTPException(404, "unknown sample")
    return FileResponse(WORLD_DIR / "docs" / files[name], filename=f"sample_{name}_{files[name]}")


# ============================================================================ documents
@app.get("/api/documents")
def documents(account_id: str | None = None, type: str | None = None, format: str | None = None, q: str | None = None):
    with LOCK:
        out = [doc_meta(d) for d in RT.store.docs.values()]
        if account_id:
            out = [d for d in out if d["account_id"] == account_id]
        if type:
            out = [d for d in out if d["doc_type"] == type]
        if format:
            out = [d for d in out if d["format"] == format]
        if q:
            ql = q.lower()
            out = [d for d in out if ql in d["title"].lower() or ql in (d["account_name"] or "").lower() or ql in d["filename"].lower()]
        return sorted(out, key=lambda d: d["received_at"], reverse=True)


def get_doc(doc_id: str) -> dict:
    d = RT.store.docs.get(doc_id) or RT.world_docs.get(doc_id)
    if not d:
        raise HTTPException(404, "document not found")
    return d


@app.get("/api/documents/{doc_id}")
def document(doc_id: str):
    with LOCK:
        return doc_meta(get_doc(doc_id))


@app.get("/api/documents/{doc_id}/raw")
def raw(doc_id: str):
    d = get_doc(doc_id)
    p = doc_path(d)
    if d["format"] in ("csv", "yaml", "txt"):
        return PlainTextResponse(p.read_text(), media_type=CT[d["format"]])
    if d["format"] in ("json", "geojson"):
        return PlainTextResponse(p.read_text(), media_type="text/plain; charset=utf-8")
    return FileResponse(p, media_type=CT.get(d["format"], "application/octet-stream"), filename=d["filename"],
                        headers={"Content-Disposition": f'inline; filename="{d["filename"]}"'})


_XLSX_CACHE: dict[str, dict] = {}


@app.get("/api/documents/{doc_id}/xlsx")
def xlsx(doc_id: str):
    d = get_doc(doc_id)
    if doc_id in _XLSX_CACHE:
        return _XLSX_CACHE[doc_id]
    wb = load_workbook(doc_path(d), data_only=False)
    sheets = []
    for ws in wb.worksheets:
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                v = c.value
                t = "n" if isinstance(v, (int, float)) and not isinstance(v, bool) else "b" if isinstance(v, bool) else "f" if isinstance(v, str) and v.startswith("=") else "s"
                cell = {"v": v if t != "f" else None, "t": t}
                if t == "f":
                    cell["f"] = v
                    # evaluate simple SUM formulas for display
                    import re as _re
                    m = _re.match(r"=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", v)
                    if m:
                        col, r0, r1 = m.group(1), int(m.group(2)), int(m.group(4))
                        cell["v"] = sum((ws[f"{col}{i}"].value or 0) for i in range(r0, r1 + 1) if isinstance(ws[f"{col}{i}"].value, (int, float)))
                if c.font and c.font.bold:
                    cell["bold"] = True
                if c.fill and c.fill.fgColor and c.fill.fill_type == "solid" and isinstance(c.fill.fgColor.rgb, str):
                    cell["fill"] = "#" + c.fill.fgColor.rgb[-6:]
                if c.number_format and c.number_format != "General":
                    cell["num_fmt"] = c.number_format
                cells[c.coordinate] = cell
        hidden_cols = [k for k, dim in ws.column_dimensions.items() if dim.hidden]
        sheets.append({"name": ws.title, "hidden": ws.sheet_state != "visible", "max_row": ws.max_row, "max_col": ws.max_column, "cells": cells,
                       "merges": [str(m) for m in ws.merged_cells.ranges], "hidden_rows": [r for r, dim in ws.row_dimensions.items() if dim.hidden],
                       "hidden_cols": hidden_cols, "col_widths": {k: dim.width for k, dim in ws.column_dimensions.items() if dim.width},
                       "freeze": ws.freeze_panes})
    _XLSX_CACHE[doc_id] = {"sheets": sheets}
    return _XLSX_CACHE[doc_id]


@app.get("/api/documents/{doc_id}/eml")
def eml(doc_id: str):
    d = get_doc(doc_id)
    msg = email.message_from_bytes(doc_path(d).read_bytes(), policy=email_policy.default)
    body = msg.get_body(preferencelist=("plain",)).get_content()
    atts = []
    linked = d.get("attachments", [])
    for i, part in enumerate(msg.iter_attachments()):
        atts.append({"filename": part.get_filename(), "doc_id": linked[i] if i < len(linked) else None, "size_bytes": len(part.get_payload(decode=True) or b""),
                     "content_type": part.get_content_type()})
    return {"from": msg["From"], "to": [x.strip() for x in (msg["To"] or "").split(",")], "cc": [], "subject": msg["Subject"], "date": d["received_at"],
            "text": body, "html": None, "attachments": atts, "highlights": RT.email_highlights.get(doc_id, [])}


@app.get("/api/documents/{doc_id}/extraction")
def extraction(doc_id: str):
    with LOCK:
        d = get_doc(doc_id)
        fields = []
        for oid in RT.store.by_doc.get(doc_id, []):
            o = RT.store.obs[oid]
            p = RT.obs_payload(o)
            fields.append({"obs_id": oid, "field_code": o.field_code, "label": p["field_label"], "subject_label": p["subject_label"],
                           "value_display": p["value_display"], "anchor": p["anchor"] or {"doc_id": doc_id, "kind": "derived"}, "confidence": p["confidence"]})
        issues = [{"code": i["code"], "label": i["label"], "anchor": ({"doc_id": doc_id, "kind": "xlsx", **i["anchor"]} if i.get("anchor") else None),
                   "severity": i["severity"]} for i in RT.store.doc_issues.get(doc_id, [])]
        return {"doc_id": doc_id, "method": RT.store.doc_method.get(doc_id, "Not extracted (reference / model output)"), "fields": fields, "issues": issues}


# ============================================================================ rules
def rule_payload(r: RE.Rule) -> dict:
    st = RT.rule_stats.get(r.rule_id, {"accepted": 0, "rejected": 0})
    fired = PR.get(r.product).rule_fired(RT, r.rule_id)
    tot = st["accepted"] + st["rejected"]
    return {"rule_id": r.rule_id, "version": r.version, "family": r.family, "title": r.title, "effective_from": r.effective_from, "effective_to": r.effective_to,
            "scope": r.scope, "applies_to": r.applies_to, "when": r.when, "outcome": r.outcome, "severity": r.severity, "expected": r.expected,
            "referral_level": r.referral_level, "impact_method": r.impact, "source": r.source or "Anaira standard library", "origin": r.origin,
            "tests": [{"name": t.get("name", ""), "fixture": t.get("fixture", {}), "expect": t["expect"]} for t in r.tests],
            "stats": {"fired": fired, "accepted": st["accepted"], "rejected": st["rejected"], "precision": (st["accepted"] / tot) if tot else None},
            "yaml": r.yaml_text, "product": r.product}


@app.get("/api/rules")
def rules(product: str | None = None):
    with LOCK:
        return [rule_payload(r) for r in RT.rules.values() if not product or r.product == product]


@app.get("/api/rules/{rid}")
def rule(rid: str):
    with LOCK:
        if rid not in RT.rules:
            raise HTTPException(404, "rule not found")
        return rule_payload(RT.rules[rid])


def parse_rule_yaml(text: str) -> RE.Rule:
    try:
        d = yaml.safe_load(text)
        if isinstance(d, list):
            d = d[0]
        return RE.rule_from_dict(d, text)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(422, f"Invalid rule: {e}")


@app.post("/api/rules/{rid}/test")
def rule_test(rid: str, b: dict = Body(...)):
    r = parse_rule_yaml(b["yaml"])
    return RE.run_tests(r)


@app.post("/api/rules/{rid}/backtest")
def rule_backtest(rid: str, b: dict = Body(...)):
    with LOCK:
        new = parse_rule_yaml(b["yaml"])
        return PR.get(new.product).backtest(RT, RT.rules.get(rid), new)


@app.post("/api/rules/{rid}/publish")
def rule_publish(rid: str, b: dict = Body(...)):
    with LOCK:
        new = parse_rule_yaml(b["yaml"])
        old = RT.rules.get(rid)
        if old and new.version <= old.version:
            new.version = old.version + 1
            d = yaml.safe_load(b["yaml"])
            d = d[0] if isinstance(d, list) else d
            d["version"] = new.version
            new.yaml_text = yaml.safe_dump(d, sort_keys=False, allow_unicode=True, width=120)
        RT.rules[new.rule_id] = new
        PR.get(new.product).reevaluate_all(RT)
        RT.log(None, "user", "03", f"Rule {new.rule_id} v{new.version} published", "Book re-evaluated", "Rule studio")
        return rule_payload(new)


# ============================================================================ data quality
@app.get("/api/data-quality")
def data_quality():
    with LOCK:
        docs = [d for d in RT.store.docs.values() if d.get("account_id")]
        n_fields = sum(len(RT.store.by_doc.get(d["doc_id"], [])) for d in docs)
        confs = [RT.store.obs[o].confidence for d in docs for o in RT.store.by_doc.get(d["doc_id"], [])]
        matching = {"matched": 0, "new": 0, "deleted": 0, "ambiguous": 0, "merged": 0, "split": 0}
        accounts = []
        queue = []
        for a, ren in RT.ren.items():
            locs = RT.store.account_locations(a)
            for e in locs.values():
                k = e.match_status.lower()
                if RT.renewal_sov.get(a) and k in matching:
                    matching[k] += 1
                if e.match_status == "AMBIGUOUS":
                    tgt = locs.get(e.proposed_target)
                    last_alias = e.aliases[-1] if e.aliases else {}
                    queue.append({"item_id": f"{a}|{e.location_uid}", "account_id": a, "account_name": RT.systems["accounts"][a]["name"], "kind": "LOCATION_MATCH",
                                  "label": f"'{last_alias.get('address', e.address)}' → {e.label}",
                                  "detail": f"{e.match_method} · score {e.match_score} · expiring address {e.address}", "options": []})
            if not ren.pass_no:
                continue
            low = sum(1 for o in RT.store.account_obs(a) if o.confidence < 0.85 and o.obs_type in ("N", "C"))
            conflicts = sum(len(s.get("_conflicts", {})) for s in ren.loc_ctx.values())
            gaps = []
            aal = {uid: s.get("aal") or 0 for uid, s in ren.loc_ctx.items()}
            tot = sum(aal.values()) or 1
            for uid, s in ren.loc_ctx.items():
                for fld, why in (("roof_year", "Missing roof year (CAT secondary modifier)"), ("year_built", "Missing year built"), ("construction_class", "Unmapped construction")):
                    if not s.get(fld) and s.get("in_current", True):
                        gaps.append({"location_label": s.get("label", ""), "field": fld, "reason": why, "impact_rank": 0, "aal_weight": aal.get(uid, 0) / tot})
            gaps.sort(key=lambda g: -g["aal_weight"])
            for i, g in enumerate(gaps, 1):
                g["impact_rank"] = i
            unmatched = sum(1 for e in locs.values() if e.match_status in ("AMBIGUOUS",))
            carried = any(s.get("_values_carried") for s in ren.loc_ctx.values()) and (ren.acct_ctx or {}).get("days_to_expiry", 999) < 90
            score = max(20, 100 - 70 * sum(g["aal_weight"] for g in gaps) - 6 * unmatched - 4 * conflicts - 0.6 * low - (12 if carried else 0))
            accounts.append({"account_id": a, "name": RT.systems["accounts"][a]["name"], "score": round(score, 1), "missing_fields": len(gaps),
                             "low_confidence": low, "stale_fields": sum(1 for s in ren.loc_ctx.values() if s.get("_values_carried")), "unmatched_locations": unmatched,
                             "conflicts": conflicts, "top_gaps": gaps[:5]})
        accounts.sort(key=lambda x: x["score"])
        return {"avg_score": round(sum(x["score"] for x in accounts) / len(accounts), 1) if accounts else 0,
                "extraction": {"documents": len(docs), "fields": n_fields, "avg_confidence": round(sum(confs) / len(confs), 3) if confs else 0,
                               "human_review_queue": len(queue)},
                "matching": matching, "accounts": accounts, "review_queue": queue}


@app.post("/api/data-quality/review/{item_id}")
def review(item_id: str, b: dict = Body(...), x_user_id: str | None = Header(None)):
    with LOCK:
        RT.match_decision(item_id, b["decision"], uid_of(x_user_id))
        return {"ok": True}


# ============================================================================ geo, pipeline, mocks, search
_HAZ: dict | None = None


@app.get("/api/geo/hazards")
def hazards():
    global _HAZ
    if _HAZ is None:
        import math
        feats = []
        def ring(lat, lon, rx, ry, n=36):
            return [[[lon + rx * math.cos(2 * math.pi * i / n) / math.cos(math.radians(lat)), lat + ry * math.sin(2 * math.pi * i / n)] for i in range(n + 1)]]
        coast = [(29.3, -94.8), (29.6, -93.5), (29.8, -91.5), (29.3, -90.0), (30.3, -88.5), (30.4, -86.8), (29.9, -85.3), (29.2, -83.2), (28.0, -82.8),
                 (26.6, -82.2), (25.8, -81.4), (25.3, -80.4), (26.2, -80.1), (27.5, -80.3), (29.0, -80.9), (30.4, -81.4), (32.0, -80.8), (32.8, -79.8),
                 (33.8, -78.5), (34.6, -76.8), (35.4, -75.5)]
        for lat, lon in coast:
            feats.append({"type": "Feature", "geometry": {"type": "Polygon", "coordinates": ring(lat, lon, 0.9, 0.55)}, "properties": {"layer": "wind", "tier": "T1", "label": "Wind Tier 1 (coastal)"}})
        for lat, lon, rx, ry, lab in ((34.0, -118.2, 1.3, 0.8, "SoCal seismic zone"), (37.7, -122.2, 0.9, 0.9, "Bay Area seismic zone"), (39.5, -119.8, 0.6, 0.6, "Walker Lane seismic zone"), (47.5, -122.3, 0.6, 0.7, "Puget Sound seismic zone")):
            feats.append({"type": "Feature", "geometry": {"type": "Polygon", "coordinates": ring(lat, lon, rx, ry)}, "properties": {"layer": "eq", "tier": "Z4", "label": lab}})
        for lat, lon, lab in ((29.76, -95.37, "Houston SFHA"), (27.95, -82.46, "Tampa Bay SFHA"), (29.95, -90.07, "New Orleans SFHA"), (32.78, -79.93, "Charleston SFHA"), (40.74, -74.17, "Newark SFHA")):
            feats.append({"type": "Feature", "geometry": {"type": "Polygon", "coordinates": ring(lat, lon, 0.25, 0.18)}, "properties": {"layer": "flood", "tier": "AE", "label": lab}})
        for lat, lon, lab in ((39.3, -120.2, "Sierra WUI"), (34.3, -117.5, "Inland Empire WUI"), (39.9, -105.3, "Front Range WUI"), (33.6, -111.6, "Phoenix WUI")):
            feats.append({"type": "Feature", "geometry": {"type": "Polygon", "coordinates": ring(lat, lon, 0.7, 0.5)}, "properties": {"layer": "wildfire", "tier": "High", "label": lab}})
        _HAZ = {"type": "FeatureCollection", "features": feats}
    return _HAZ


PIPE = PR.get("renewal").PIPE


@app.get("/api/products")
def products():
    return PR.metas()


@app.get("/api/pipeline")
def pipeline(product: str = "renewal"):
    m = PR.get(product)
    with LOCK:
        counts = m.pipeline_counts(RT)
        return [{"code": c, "name": n, "group": g, "component": comp, "mode": mo, "description": desc,
                 "counts": [{"label": k, "value": v} for k, v in counts.get(c, [])], "port": port, "product": m.ID}
                for c, n, g, comp, mo, desc, port in m.PIPE]


@app.get("/api/pipeline-groups")
def pipeline_groups(product: str = "renewal"):
    return PR.get(product).GROUPS


@app.get("/api/pipeline-subjects")
def pipeline_subjects(product: str = "renewal"):
    with LOCK:
        return PR.get(product).subjects(RT)


@app.get("/api/pipeline/journey/{acct}")
def journey(acct: str, product: str = "renewal"):
    m = PR.get(product)
    with LOCK:
        out = []
        for c, *_ in m.PIPE:
            v = m.build_stage(RT, c, acct)
            out.append({"code": c, "name": v["name"], "mode": v["mode"], "status": v["status"], "headline": v["headline"], "has_action": bool(v["next_action"])})
        return out


@app.get("/api/pipeline/{code}")
def stage(code: str, account_id: str | None = None, product: str = "renewal"):
    m = PR.get(product)
    with LOCK:
        if code not in {p[0] for p in m.PIPE}:
            raise HTTPException(404, "unknown stage")
        return {**m.build_stage(RT, code, account_id or m.subjects(RT)[0]["id"]), "product": m.ID}


@app.post("/api/pipeline/{code}/run")
def stage_run(code: str, b: dict = Body(...)):
    m = PR.get(b.get("product"))
    with LOCK:
        acct = b.get("account_id") or m.subjects(RT)[0]["id"]
        try:
            msg = m.run_stage(RT, code, acct)
        except (ValueError, PermissionError) as e:
            raise HTTPException(409, str(e))
        return {"message": msg, "stage": {**m.build_stage(RT, code, acct), "product": m.ID}}


def _pb_state(pid: str) -> dict:
    st = getattr(RT, "playbook_state", None)
    if st is None:
        RT.playbook_state = st = {}
    return st.setdefault(pid, {"progress": 0, "results": []})


def _brief(p: dict) -> dict:
    m = PR.get(p.get("product"))
    b = dict(p.get("brief") or {})
    b.setdefault("real", m.REAL_CORE)
    b.setdefault("mocked", m.MOCK_CORE)
    return b


def _pb_payload(p: dict) -> dict:
    from uwc.narration import for_playbook
    m = PR.get(p.get("product"))
    st = _pb_state(p["id"])
    nar = for_playbook(p)
    return {**{k: p.get(k) for k in ("id", "title", "account_id", "account_name", "scenario", "aspects")}, "product": m.ID,
            "subject_href": p.get("subject_href") or (f"/accounts/{p['account_id']}" if m.ID == "renewal" else None),
            "intro": nar["intro"], "outro": nar["outro"], "brief": _brief(p), "facts": m.facts(RT, p["account_id"]),
            "steps": [{"index": i, "code": s["code"], "label": s["label"], "action": s["action"], "say": nar["steps"][i] if i < len(nar["steps"]) else s["label"],
                       "status": "DONE" if i < st["progress"] else "CURRENT" if i == st["progress"] else "UPCOMING",
                       "result": st["results"][i] if i < len(st["results"]) else None,
                       "then_say": PB.speakable(st["results"][i]) if i < len(st["results"]) and s["action"] != "view" else None} for i, s in enumerate(p["steps"])],
            "progress": st["progress"], "done": st["progress"] >= len(p["steps"]), "clock": RT.clock}


@app.get("/api/playbooks")
def playbooks(product: str | None = None):
    with LOCK:
        return [_pb_payload(p) for p in PR.all_playbooks(RT) if not product or p["product"] == product]


def _find_pb(pid: str) -> dict:
    try:
        return PR.find_playbook(RT, pid)
    except KeyError:
        raise HTTPException(404, "playbook not found")


@app.get("/api/playbooks/{pid}")
def playbook(pid: str):
    with LOCK:
        return _pb_payload(_find_pb(pid))


@app.post("/api/playbooks/{pid}/start")
def playbook_start(pid: str, b: dict = Body(default={})):
    with LOCK:
        p = _find_pb(pid)
        if b.get("reset", True):
            RT.reset()
        RT.playbook_state = {}
        return _pb_payload(p)


@app.post("/api/playbooks/{pid}/step")
def playbook_step(pid: str):
    with LOCK:
        p = _find_pb(pid)
        m = PR.get(p["product"])
        st = _pb_state(pid)
        i = st["progress"]
        if i >= len(p["steps"]):
            return {**_pb_payload(p), "message": "Playbook complete"}
        try:
            msg = m.run_step(RT, p, p["steps"][i])
        except (ValueError, PermissionError, KeyError) as e:
            msg = f"Step could not run: {e}"
            st["results"].append(msg)
            raise HTTPException(409, msg)
        st["results"].append(msg)
        st["progress"] = i + 1
        return {**_pb_payload(p), "message": msg}


@app.get("/api/mocks/systems")
def mocks(product: str = "renewal"):
    with LOCK:
        return PR.get(product).mocks(RT)


@app.get("/api/mocks/outbox")
def outbox(product: str | None = None):
    with LOCK:
        return [m for m in reversed(RT.outbox) if not product or m.get("product", "renewal") == product]


@app.get("/api/search")
def search(q: str = Query("")):
    with LOCK:
        ql = q.lower().strip()
        hits = []
        for a, reg in RT.systems["accounts"].items():
            if ql in reg["name"].lower() or (reg.get("scenario") or "").lower() == ql:
                hits.append({"kind": "account", "id": a, "title": reg["name"], "subtitle": reg.get("scenario_title") or reg["broker"], "href": f"/accounts/{a}"})
        for d in RT.store.docs.values():
            if ql in d["title"].lower() or ql in d["filename"].lower():
                hits.append({"kind": "document", "id": d["doc_id"], "title": d["title"], "subtitle": d.get("account_name") or d["doc_type"], "href": f"/documents/{d['doc_id']}"})
        for r in RT.rules.values():
            if ql in r.rule_id.lower() or ql in r.title.lower():
                hits.append({"kind": "rule", "id": r.rule_id, "title": r.title, "subtitle": r.rule_id, "href": f"/rules/{r.rule_id}"})
        return hits[:20]


@app.exception_handler(Exception)
async def err(_, exc: Exception):
    import traceback
    traceback.print_exc()
    return JSONResponse({"detail": f"{type(exc).__name__}: {exc}"}, status_code=500)


# ============================================================================ product-specific routers
for _m in PR.MODULES.values():
    if getattr(_m, "router", None) is not None:
        app.include_router(_m.router)
