"""Render carrier, broker and vendor documents from truth."""
from __future__ import annotations

import random
from email.message import EmailMessage
from email.utils import format_datetime
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from uwc.config import CARRIER
from uwc.refdata import FORMS, GUIDELINES, SAFEGUARD_CODES, AUTHORITY, CONSTRUCTION, OCCUPANCY
from uwc.world.pdfkit import Pdf, money, pctf
from uwc.world.truth import Account, Email, Location, QuoteTruth, Terms

# --------------------------------------------------------------------- SOV
TEMPLATES = {
    "A": ["Loc #", "Bldg #", "Location Name", "Street Address", "City", "State", "Zip", "Occupancy", "Construction", "Year Built",
          "# Stories", "Sq Ft", "Roof Year", "Sprinklered", "Building Value", "Contents / BPP", "Stock", "Business Income", "TIV", "Valuation"],
    "B": ["Loc No", "Site Name", "Address", "City", "ST", "ZIP", "Occupancy Description", "Const Type", "Yr Blt", "Stories", "Area (sf)",
          "Roof Updated", "Sprinkler %", "Bldg RC", "BPP", "Inventory", "BI/EE (12 mo)", "Total Insured Value"],
    "C": ["Site", "Building", "Site Description", "Location Address", "City", "State", "Postal Code", "Use", "ISO Class / Construction",
          "Year of Construction", "No. of Floors", "Gross Area", "Roof Replaced", "Sprinkler Protection", "Building (RC)", "Contents",
          "Stock / Inventory", "Time Element", "Total"],
}
BROKER_TEMPLATE = {"Harlan & Pierce Risk Partners": "A", "Keel & Crane Insurance Services": "B", "Brightwater Risk Advisors": "C",
                   "Calder Street Brokerage": "A", "Ironbridge Specialty Brokers": "B", "Tidewater Commercial Insurance": "C"}

HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="C9D1DC")


def _spr(l: Location, tmpl: str):
    if tmpl == "B":
        return round(l.sprinkler_pct * 100)
    if tmpl == "C":
        return "Full" if l.sprinkler_pct >= 0.99 else "Partial" if l.sprinkler_pct > 0 else "None"
    return "Y" if l.sprinkler_pct >= 0.99 else "Partial" if l.sprinkler_pct > 0 else "N"


def render_sov(a: Account, which: str, path: Path, as_of: str) -> dict:
    """Render prior ('2025') or current ('2026') SOV. Returns render notes (for the regression record)."""
    tmpl = BROKER_TEMPLATE.get(a.broker, "A")
    heads = list(TEMPLATES[tmpl])
    q = a.sov_quirks
    scale = 1000.0 if (which == "current" and q.get("scale_000s")) else 1.0
    has_storage = any(l.storage_height_ft for l in a.locations)
    if has_storage:
        heads.insert(heads.index(heads[-1]), "Max Storage Ht (ft)")
    wb = Workbook()
    ws = wb.active
    ws.title = "SOV" if tmpl != "C" else "Locations"
    ws["A1"] = f"{a.name} — Statement of Values"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    ws["A2"] = f"Values as of {as_of}" + ("   (values in $000s)" if scale == 1000 else "")
    ws["A2"].font = Font(italic=True, size=10, color="555555")
    ws["A3"] = f"Prepared by {a.broker}"
    ws["A3"].font = Font(size=9, color="777777")
    hr = q.get("header_row", 5) if which == "current" else 5
    for ci, h in enumerate(heads, 1):
        c = ws.cell(row=hr, column=ci, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=THIN)
        ws.column_dimensions[get_column_letter(ci)].width = max(9, min(28, len(h) + 4))
    ws.row_dimensions[hr].height = 30
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    r = hr + 1
    notes: dict = {"template": tmpl, "header_row": hr, "scale": scale, "rows": {}}
    locs = [l for l in a.locations if (l.values_prior if which == "prior" else l.values_current)]
    if which == "current" and q.get("renumbered"):
        locs = sorted(locs, key=lambda l: l.loc_no_current or "")
    first_data = r
    for l in locs:
        v = l.values_prior if which == "prior" else l.values_current
        loc_no = l.loc_no_prior if which == "prior" else l.loc_no_current
        occ_raw = (l.occupancy_raw_prior or l.occupancy_raw) if which == "prior" else l.occupancy_raw
        roof = l.roof_year if l.status == "new" else (l.sov_roof_year or l.roof_year)
        if q.get("drop_roof_year") and which == "current":
            roof = None
        addr = l.address
        if which == "current" and l.sov_address_variant:
            addr = l.sov_address_variant
        cons_raw = l.sov_construction_raw or l.construction_raw
        sqft = l.sqft
        if which == "prior" and l.sqft_prior:
            sqft = l.sqft_prior
        tiv = v.tiv
        vals = {"bldg": v.building / scale, "cont": v.contents / scale, "stock": v.stock / scale, "bi": v.bi / scale, "tiv": tiv / scale}
        row_by_head = {
            "Loc #": loc_no, "Loc No": loc_no, "Site": loc_no, "Bldg #": 1, "Building": 1,
            "Location Name": l.name, "Site Name": l.name, "Site Description": l.name,
            "Street Address": addr, "Address": addr, "Location Address": addr,
            "City": l.city, "State": l.state, "ST": l.state, "Zip": l.zip, "ZIP": l.zip, "Postal Code": l.zip,
            "Occupancy": occ_raw, "Occupancy Description": occ_raw, "Use": occ_raw,
            "Construction": cons_raw, "Const Type": cons_raw, "ISO Class / Construction": cons_raw,
            "Year Built": l.year_built, "Yr Blt": l.year_built, "Year of Construction": l.year_built,
            "# Stories": l.stories, "Stories": l.stories, "No. of Floors": l.stories,
            "Sq Ft": sqft, "Area (sf)": sqft, "Gross Area": sqft,
            "Roof Year": roof, "Roof Updated": roof, "Roof Replaced": roof,
            "Sprinklered": _spr(l, tmpl), "Sprinkler %": _spr(l, tmpl), "Sprinkler Protection": _spr(l, tmpl),
            "Building Value": vals["bldg"], "Bldg RC": vals["bldg"], "Building (RC)": vals["bldg"],
            "Contents / BPP": vals["cont"], "BPP": vals["cont"], "Contents": vals["cont"],
            "Stock": vals["stock"], "Inventory": vals["stock"], "Stock / Inventory": vals["stock"],
            "Business Income": vals["bi"], "BI/EE (12 mo)": vals["bi"], "Time Element": vals["bi"],
            "TIV": vals["tiv"], "Total Insured Value": vals["tiv"], "Total": vals["tiv"],
            "Valuation": "RC",
            "Max Storage Ht (ft)": ((l.sprinkler_design_ft if l.occupancy_prior else l.storage_height_ft) if which == "prior" else l.storage_height_ft),
        }
        for ci, h in enumerate(heads, 1):
            val = row_by_head.get(h)
            c = ws.cell(row=r, column=ci, value=val)
            if isinstance(val, (int, float)) and h not in ("Loc #", "Loc No", "Site", "Bldg #", "Building", "Year Built", "Yr Blt", "Year of Construction", "Roof Year", "Roof Updated", "Roof Replaced", "# Stories", "Stories", "No. of Floors", "Sprinkler %", "Max Storage Ht (ft)"):
                c.number_format = "#,##0"
        notes["rows"][l.key] = r
        r += 1
    last_data = r - 1
    if (which == "current" and q.get("totals_row")) or (which == "prior" and len(locs) > 6):
        ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
        for ci, h in enumerate(heads, 1):
            if h in ("Building Value", "Bldg RC", "Building (RC)", "Contents / BPP", "BPP", "Contents", "Stock", "Inventory", "Stock / Inventory",
                     "Business Income", "BI/EE (12 mo)", "Time Element", "TIV", "Total Insured Value", "Total"):
                col = get_column_letter(ci)
                c = ws.cell(row=r, column=ci, value=f"=SUM({col}{first_data}:{col}{last_data})")
                c.number_format = "#,##0"
                c.font = Font(bold=True)
        notes["totals_row"] = r
    # a hidden internal-notes column (common broker artefact)
    nc = len(heads) + 1
    ws.cell(row=hr, column=nc, value="Internal notes").font = Font(bold=True, color="999999")
    ws.column_dimensions[get_column_letter(nc)].hidden = True
    notes["hidden_col"] = get_column_letter(nc)
    # instructions sheet (hidden)
    s2 = wb.create_sheet("Broker notes")
    s2["A1"] = "Values supplied by insured; broker has not verified. Replacement cost basis unless noted."
    s2.sheet_state = "hidden"
    wb.save(path)
    return notes


# --------------------------------------------------------------------- terms helpers
def _terms_pairs(t: Terms) -> list[tuple[str, str]]:
    return [
        ("Limit of insurance", money(t.limit)), ("Limit basis", {"blanket": "Blanket", "scheduled": "Scheduled", "loss_limit": "Loss limit"}[t.limit_basis]),
        ("AOP deductible", money(t.aop_deductible)), ("Named storm deductible", pctf(t.named_storm_ded_pct) + (" per location" if t.named_storm_ded_pct else "")),
        ("Named storm minimum", money(t.named_storm_ded_min)), ("Wind/hail deductible", pctf(t.wind_hail_ded_pct)),
        ("Business income sublimit", money(t.bi_sublimit) if t.bi_sublimit else "Included in limit"),
        ("Flood sublimit", money(t.flood_sublimit) if t.flood_sublimit else "Not covered" if t.flood_sublimit is None else money(t.flood_sublimit)),
        ("Earthquake sublimit", money(t.eq_sublimit) if t.eq_sublimit else "Not covered"),
    ]


def _forms_table(p: Pdf, t: Terms):
    p.section("Forms and endorsements")
    p.table(["Form", "Edition", "Title"], [(f, "10 12" if not f.startswith("NS") else "01 25", FORMS.get(f, "")) for f in t.forms], [90, 60, 350])


def _safeguards_table(p: Pdf, a: Account, t: Terms):
    if not t.safeguards:
        return
    p.section("Protective safeguards schedule (CP 04 11)")
    rows = []
    for sg in t.safeguards:
        l = a.loc(sg["location_key"])
        rows.append((l.loc_no_prior or l.loc_no_current or "", l.address, sg["code"], SAFEGUARD_CODES[sg["code"]]))
    p.table(["Loc", "Address", "Symbol", "Protective device or service"], rows, [40, 190, 50, 220])


def _loc_schedule(p: Pdf, a: Account, prior: bool = True):
    p.section("Schedule of locations")
    rows = []
    for l in a.locations:
        v = l.values_prior if prior else l.values_current
        if not v:
            continue
        rows.append((l.loc_no_prior if prior else l.loc_no_current, l.address, l.city, l.state, OCCUPANCY[l.occupancy_prior or l.occupancy][0] if prior else OCCUPANCY[l.occupancy][0],
                     money(v.building), money(v.contents + v.stock), money(v.bi), money(v.tiv)))
    p.table(["Loc", "Address", "City", "ST", "Occupancy", "Building", "BPP/Stock", "BI", "TIV"], rows,
            [28, 110, 60, 24, 100, 62, 62, 52, 66], size=7.2, align_right={5, 6, 7, 8})


def render_declarations(a: Account, path: Path, issue_date: str, terms: Terms | None = None, premium: float | None = None,
                        policy_no: str | None = None, term: tuple[str, str] | None = None, prior: bool = True):
    t = terms or a.terms_issued
    p = Pdf(path, "carrier", "Commercial Property Declarations", policy_no or a.policy_no)
    p.title("Commercial Property Declarations", f"Issued {issue_date} · {CARRIER}")
    ts, te = term or (a.term_start, a.term_end)
    p.kv([("Policy number", policy_no or a.policy_no), ("Named insured", a.name), ("Policy period", f"{ts} to {te}"),
          ("Producer", a.broker), ("Carrier share", f"{a.carrier_share * 100:.0f}%"),
          ("Layer", "Primary" if not a.layer_attach else f"{money(a.layer_limit)} xs {money(a.layer_attach)}"),
          ("Total premium", money(premium if premium is not None else a.premium_expiring)), ("Admitted", "Yes" if a.admitted else "No — surplus lines")])
    p.section("Limits and deductibles")
    p.kv(_terms_pairs(t))
    _loc_schedule(p, a, prior=prior)
    _forms_table(p, t)
    _safeguards_table(p, a, t)
    p.section("Conditions")
    p.para("This policy is subject to the Commercial Property Conditions and Common Policy Conditions. Protective safeguards scheduled above "
           "must be maintained in complete working order; the insured must notify the company of any impairment. Coverage for a building "
           "vacant beyond 60 consecutive days is modified per the Commercial Property Conditions unless a vacancy permit is endorsed.", size=8.5)
    p.signature("Authorized representative", CARRIER, issue_date)
    p.save()


def render_binder(a: Account, path: Path, bind_date: str, terms: Terms | None = None, premium: float | None = None,
                  subjectivities: list[dict] | None = None, binder_no: str | None = None, term: tuple[str, str] | None = None, prior: bool = True):
    t = terms or a.terms_binder
    bn = binder_no or a.policy_no.replace("NSP", "BND")
    p = Pdf(path, "carrier", "Binder of Insurance", bn)
    p.title("Binder of Insurance", f"Bound {bind_date}")
    ts, te = term or (a.term_start, a.term_end)
    p.kv([("Binder number", bn), ("Named insured", a.name), ("Policy period", f"{ts} to {te}"), ("Producer", a.broker),
          ("Bound premium", money(premium if premium is not None else a.premium_expiring)), ("Carrier share", f"{a.carrier_share * 100:.0f}%")])
    p.section("Bound terms")
    p.kv(_terms_pairs(t))
    _forms_table(p, t)
    _safeguards_table(p, a, t)
    subs = subjectivities if subjectivities is not None else a.subjectivities
    if subs:
        p.section("Subjectivities")
        p.table(["Subjectivity", "Due", "Status at bind"], [(s["text"], s["due"], "Outstanding") for s in subs], [330, 80, 90])
    p.para("This binder is issued subject to the terms of the policy to be issued and the subjectivities above. It terminates on issuance of the policy.", size=8.5)
    p.signature("Underwriter", CARRIER, bind_date)
    p.save()


def render_quote(a: Account, q: QuoteTruth, path: Path, term: tuple[str, str] | None = None, subjectivities: list[dict] | None = None, created_by: str = ""):
    p = Pdf(path, "carrier", f"Property Quotation v{q.version}", a.policy_no.replace("NSP", "QTE") + f"-v{q.version}")
    p.title(f"Property Quotation — version {q.version}", f"Issued {q.date}" + (f" · prepared by {created_by}" if created_by else ""))
    ts, te = term or (a.term_start, a.term_end)
    p.kv([("Quote version", str(q.version)), ("Named insured", a.name), ("Proposed period", f"{ts} to {te}"), ("Producer", a.broker),
          ("Quoted premium", money(q.premium)), ("Valid until", q.date[:8] + "28")])
    p.section("Quoted terms")
    p.kv(_terms_pairs(q.terms))
    _forms_table(p, q.terms)
    subs = subjectivities if subjectivities is not None else a.subjectivities
    if subs:
        p.section("Subjectivities")
        p.table(["Subjectivity", "Due"], [(s["text"], s["due"]) for s in subs], [400, 100])
    p.save()


def render_endorsement(a: Account, e, path: Path):
    p = Pdf(path, "carrier", "Policy Change Endorsement", e.endt_id)
    p.title("Policy Change Endorsement", f"Effective {e.effective}")
    p.kv([("Endorsement number", e.endt_id), ("Policy number", a.policy_no), ("Named insured", a.name), ("Effective date", e.effective),
          ("Change type", e.type), ("Additional / return premium", money(e.premium_delta))])
    p.section("Description of change")
    p.para(e.description)
    p.save()


def render_loss_run(a: Account, as_of: str, path: Path):
    p = Pdf(path, "carrier", "Loss Run", a.policy_no)
    p.title("Loss Run — Commercial Property", f"Valued as of {as_of} · 5 policy years")
    tot_paid = sum(c.paid for c in a.claims)
    tot_res = sum(c.reserve for c in a.claims)
    p.kv([("Named insured", a.name), ("Policy number", a.policy_no), ("Valuation date", as_of), ("Claim count", str(len(a.claims))),
          ("Total paid", money(tot_paid)), ("Total incurred", money(tot_paid + tot_res))])
    rows = []
    for c in sorted(a.claims, key=lambda c: c.dol):
        loc = a.loc(c.location_key).address if c.location_key else "—"
        rows.append((c.claim_id, c.dol, loc, c.cause.replace("_", " "), c.status, money(c.paid), money(c.reserve), money(c.paid + c.reserve)))
    if not rows:
        p.para("No losses reported during the experience period.")
    else:
        p.table(["Claim #", "Date of loss", "Location", "Cause", "Status", "Paid", "Reserve", "Incurred"], rows, [70, 60, 120, 70, 45, 55, 55, 60], align_right={5, 6, 7})
        p.section("Loss descriptions")
        for c in sorted(a.claims, key=lambda c: c.dol):
            p.para(f"{c.claim_id} — {c.description}", size=8.2)
    p.save()


def render_engineering(a: Account, survey_date: str, locs: list[Location], path: Path, engineer: str = "Elena Brooks, CSP"):
    p = Pdf(path, "engineering", "Property Risk Engineering Report", f"{a.account_id.upper()}-{survey_date}")
    p.title("Property Risk Engineering Report", f"Survey date {survey_date} · {engineer}")
    p.kv([("Insured", a.name), ("Survey date", survey_date), ("Engineer", engineer), ("Scope", "COPE, fire protection, natural hazard, recommendations")], cols=2)
    for l in locs:
        p.section(f"Location {l.loc_no_prior or ''} — {l.name}")
        occ = l.occupancy_prior or l.occupancy
        p.kv([("Address", f"{l.address}, {l.city}, {l.state}")], cols=1, label_w=150)
        pairs = [("Construction", CONSTRUCTION[l.construction][0]),
                 ("Year built", str(l.year_built)), ("Stories", str(l.stories)), ("Floor area (sq ft)", f"{l.sqft:,}"),
                 ("Roof covering", l.roof_type), ("Roof year (verified)", str(l.roof_year)),
                 ("Occupancy observed", OCCUPANCY[occ][0]), ("Sprinkler protection", f"{l.sprinkler_pct * 100:.0f}% of area"),
                 ("Fire alarm", "Central station" if l.fire_alarm else "None"), ("Public protection class", str(l.ppc))]
        if l.storage_height_ft:
            design = l.sprinkler_design_ft or 20
            observed = min(l.storage_height_ft, design) if l.occupancy_prior else l.storage_height_ft
            pairs += [("Max storage height (ft)", str(int(observed))), ("Sprinkler design storage height (ft)", str(int(design))),
                      ("Commodity", "General merchandise (Class III)" if l.occupancy_prior else (l.commodity or ""))]
        if l.cooking:
            pairs += [("Cooking suppression", "UL 300 wet chemical")]
        p.kv(pairs, cols=2, label_w=150)
        p.para(_eng_narrative(l), size=8.5)
    recs = [r for r in a.recs if r.raised == survey_date]
    if recs:
        p.section("Recommendations")
        p.table(["Rec #", "Loc", "Category", "Recommendation", "Priority", "Target date"],
                [(r.rec_id, a.loc(r.location_key).loc_no_prior or "", r.category, r.description, r.severity.title(), r.due) for r in recs],
                [44, 26, 80, 250, 50, 60], size=7.2)
        if any(r.bind_condition for r in recs):
            p.para("Recommendations marked CRITICAL are a condition of binding and must be completed and verified by Northgate Risk Engineering.", size=8.5, color="#8a1c1c")
    p.signature(engineer, "Northgate Risk Engineering", survey_date)
    p.save()


def _eng_narrative(l: Location) -> str:
    base = (f"The {l.name} facility is a {l.stories}-story {CONSTRUCTION[l.construction][0].split(' (')[0].lower()} building of approximately "
            f"{l.sqft:,} sq ft built in {l.year_built}. Housekeeping was observed to be generally good. ")
    if l.sprinkler_pct >= 0.99:
        base += "The building is fully protected by automatic sprinklers supplied from the municipal main; the system was last tested within 12 months. "
    elif l.sprinkler_pct > 0:
        base += f"Sprinkler protection covers about {l.sprinkler_pct * 100:.0f}% of the floor area; the unprotected area is used for finishing and storage. "
    else:
        base += "The building is not sprinklered. "
    if l.storage_height_ft:
        base += "Storage is in selective steel racking; storage heights were within the sprinkler design at the time of survey. "
    return base


def render_appraisal(a: Account, path: Path, date: str):
    p = Pdf(path, "appraisal", "Insurable Value Appraisal", f"CVS-{a.account_id.upper()}")
    p.title("Insurable Value Appraisal", f"Effective date {date} · Replacement cost new")
    p.kv([("Client", a.name), ("Effective date", date), ("Valuation basis", "Replacement cost new"), ("Method", "Unit-in-place, local cost modifiers")])
    rows = []
    for l in a.locations:
        rc = l.values_prior.building * 1.0 if l.values_prior else 0
        rows.append((l.name, f"{l.sqft:,}", money(rc), f"${rc / l.sqft:,.0f}"))
    p.table(["Property", "Area (sf)", "Replacement cost", "Cost / sf"], rows, [220, 80, 110, 90], align_right={1, 2, 3})
    p.para("Values reflect construction costs as of the effective date and should be updated at least every three years or after material changes.", size=8.5)
    p.signature("Crestview Valuation Services", "Senior Appraiser", date)
    p.save()


def render_alarm_certs(a: Account, path: Path, date: str, rng: random.Random) -> dict:
    p = Pdf(path, "alarm", "Burglar Alarm Certificates", a.name)
    p.title("Burglar Alarm System Certificates", f"Compiled {date} · {a.name}")
    present = {}
    for l in a.locations:
        if l.key in ("L9", "L12"):
            continue  # certificates not provided
        present[l.key] = True
        p.section(f"Certificate — {l.name}")
        p.kv([("Protected premises", f"{l.address}, {l.city}, {l.state}"), ("Alarm type", "Intrusion — motion, contact, glass-break"),
              ("Central station monitoring", "Yes — UL listed central station" if l.burglar_monitored else "No — local alarm only"),
              ("Certificate date", date), ("Last test", "2026-05-" + str(rng.randint(10, 28)))], cols=1, label_w=170)
    p.save()
    return present


def render_hood_certs(a: Account, path: Path, date: str) -> dict:
    p = Pdf(path, "hood", "Kitchen Exhaust Cleaning Certificates", a.name)
    p.title("Kitchen Exhaust Hood & Duct Cleaning Certificates", f"Compiled {date}")
    present = {}
    for l in a.locations:
        if not l.hood_cleaning_ok:
            continue
        present[l.key] = True
        p.section(f"{l.name}")
        p.kv([("Premises", f"{l.address}, {l.city}, {l.state}"), ("Service", "Hood, duct & fan cleaned to bare metal (NFPA 96)"),
              ("Service date", "2026-06-12"), ("Suppression system inspected", "Yes — UL 300 wet chemical, tagged")], cols=1, label_w=170)
    p.save()
    return present


def render_impairment_notice(a: Account, path: Path, date: str):
    l = a.locations[0]
    p = Pdf(path, "fire", "Fire Protection Impairment Notice", f"FPB-{date}")
    p.title("Notice of Fire Protection System Impairment", f"Issued {date}")
    p.kv([("Premises", f"{l.name}, {l.address}, {l.city}, {l.state}"), ("System", "Automatic sprinkler — anchor building riser"),
          ("Impairment", "System shut off; water supply valve closed during vacant-unit renovation"), ("Date observed", date),
          ("Owner notified", "Yes — property manager"), ("Expected restoration", "Unknown")], cols=1, label_w=150)
    p.para("NFPA 25 impairment procedures were not followed: no fire watch, no notification to the insurer. The owner is directed to restore the system.")
    p.save()


def render_guidelines(version: str, path: Path) -> dict[str, int]:
    g = GUIDELINES[version]
    p = Pdf(path, "carrier", g["doc_title"], f"Version {g['version']}")
    p.title(g["doc_title"], f"Version {g['version']} · effective {g['effective_from']}")
    p.para("These guidelines govern commercial property underwriting for Northgate Specialty Insurance Co. Deviations require a documented "
           "referral to the authority level shown in §3.1. Guideline parameters are machine-readable in the Northgate rule library.")
    sections = [
        ("§2.1", "Appetite", f"Target classes: office, retail, hospitality, light manufacturing, warehouse, healthcare, education, data centers. "
                            f"Declined classes: {', '.join(g['declined_classes']) or 'none'} (scrap metal and recycling yards are declined effective v2026). "
                            "Existing accounts in declined classes are to be non-renewed with statutory notice or referred for exception."),
        ("§3.1", "Authority", "Level 1: account TIV to $200M, price deviation ±5%, renewal rate change (RARC) not below −5%. Level 2: TIV to $400M, ±10%, RARC ≥ −10%. "
                             "Level 3: TIV to $1.5B, ±15%, RARC ≥ −20%, guideline exceptions. Level 4 (CUO): all others. Approvals apply to the specific terms approved."),
        ("§4.2", "Price adequacy", f"Proposed premium must be at least {g['adequacy_floor'] * 100:.0f}% of the technical premium produced by the current rating model."),
        ("§5.1", "Valuation", f"Reported building values below {g['valuation_ratio_min'] * 100:.0f}% of modelled replacement cost require an updated appraisal or a valuation debit. "
                             f"Appraisals older than {g['appraisal_max_age_years']} years are stale."),
        ("§6.2", "Roofs", f"Roofs older than {g['roof_age_max']} years in hail or wind zones require inspection or a roof schedule."),
        ("§6.5", "Sprinklers", "Storage heights above the sprinkler design basis are a critical deficiency; bind only with a re-survey condition. "
                              "Lithium-ion battery storage requires specific engineering review."),
        ("§6.8", "Commercial cooking", "Commercial cooking operations require UL 300 suppression and semi-annual hood cleaning evidenced by certificates."),
        ("§7.3", "Named storm", f"Tier 1 wind locations: minimum named storm deductible {g['wind_ded_floor_pct_t1'] * 100:.0f}% per location, "
                               f"minimum ${g['ns_min_floor']:,}. Tier 2: {g['wind_ded_floor_pct_t2'] * 100:.0f}%."),
        ("§8.4", "Security", f"High-value stock above ${g['high_value_stock']:,} in locations with a burglary score at or above {g['burglary_threshold']} "
                            "requires a UL-certificated centrally monitored intrusion alarm."),
        ("§9.2", "Vacancy", f"Buildings more than 50% vacant for over {g['vacancy_days_max']} days require referral; vacancy with impaired sprinklers is critical."),
        ("§10.1", "Engineering", "Critical recommendations must be verified closed by Northgate Risk Engineering. Bind conditions not verified are a broken commitment."),
        ("§11.4", "Accumulation", f"Renewals that take a CAT zone above {g['accumulation_util_refer'] * 100:.0f}% of its 1-in-250 capacity are referred to Level 3."),
    ]
    pages = {}
    for code, title, text in sections:
        p.new_page()
        pages[code] = p.page
        p.section(f"{code} {title}")
        p.para(text)
        p.para("Underwriters must document the evidence relied on and any deviation in the underwriting file. Where evidence conflicts, "
               "verified engineering observations take precedence over broker-reported data.", size=8.5, color="#555555")
    p.save()
    return pages


def render_authority_matrix(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Authority"
    heads = ["Level", "Roles", "Max account TIV", "Max location TIV", "Price deviation", "Min RARC", "Max premium", "Tier-1 CAT TIV", "Guideline exceptions"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        ws.column_dimensions[get_column_letter(i)].width = 18
    roles = {1: "Property Underwriter", 2: "Senior Property Underwriter", 3: "Head of Property", 4: "CUO"}
    for lvl, a in AUTHORITY.items():
        ws.append([lvl, roles[lvl], a["max_account_tiv"], a["max_loc_tiv"], a["max_price_dev"], a["min_rarc"], a["max_premium"], a["cat_t1_tiv"], "Yes" if a["can_override_guideline"] else "No"])
        for col in (3, 4, 7, 8):
            ws.cell(row=ws.max_row, column=col).number_format = "$#,##0"
        for col in (5, 6):
            ws.cell(row=ws.max_row, column=col).number_format = "0%"
    wb.save(path)


def render_eml(e: Email, attachments: list[tuple[str, Path, str]], path: Path):
    m = EmailMessage()
    m["From"] = e.sender
    m["To"] = ", ".join(e.to)
    m["Subject"] = e.subject
    m["Date"] = format_datetime(datetime.fromisoformat(e.date + "T09:14:00-04:00"))
    import hashlib
    m["Message-ID"] = f"<{hashlib.md5((e.subject + e.date).encode()).hexdigest()[:16]}@broker.example>"
    m.set_content(e.body)
    for fname, p, ctype in attachments:
        maintype, subtype = ctype.split("/")
        m.add_attachment(p.read_bytes(), maintype=maintype, subtype=subtype, filename=fname)
    path.write_bytes(bytes(m))
