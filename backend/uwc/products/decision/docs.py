"""Render the new-business submission documents from a case's truth (ACORD-style application, SOV, loss runs,
loss-control inspection, manuscript wording, broker email) and the documents produced while the demo runs
(broker replies, inspection re-survey, quote and binder). Layouts follow the shared Pdf kit so the same
extractor reads them back; nothing the engine uses is taken from the truth directly."""
from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter

from uwc.config import CARRIER
from uwc.refdata import CITIES, CONSTRUCTION, OCCUPANCY, USER_BY_ID
from uwc.world.pdfkit import Pdf, money
from uwc.world.render_docs import BROKER_TEMPLATE, HDR_FILL, HDR_FONT, TEMPLATES, THIN

from .model import NB_AUTHORITY, NB_SECTIONS, NB_STANDARDS
from .world import tiv_of


def _v(loc: dict, src: str, key: str):
    """Value of a location attribute as a given source states it (broker claims override reality)."""
    return loc.get("claims", {}).get(src, {}).get(key, loc.get(key))


def _spr_text(p: float) -> str:
    return "Yes — 100%" if p >= 0.99 else f"Partial — {p * 100:.0f}%" if p > 0 else "No"


def loc_address(loc: dict, stt: str) -> str:
    return f"{loc['address']}, {loc['city']}, {stt} {loc['zip']}"


def state_of(loc: dict) -> str:
    return CITIES[loc["city"]][2]


def _email_of(name: str, broker: str) -> str:
    dom = "".join(ch for ch in broker.split()[0].lower() if ch.isalpha())
    return f"{name.split()[0].lower()}.{name.split()[-1].lower()}@{dom}.example"


# ============================================================================ ACORD 125 / 140 style application
def render_application(c: dict, path: Path, signed: str):
    p = Pdf(path, "broker", "Commercial Insurance Application", f"ACORD 125 / 140 format · {c['insured']}")
    p.title("Commercial Insurance Application", f"ACORD 125 (applicant) and 140 (property section) format · signed {signed}")
    p.section("Applicant information")
    p.kv([("Named insured", c["insured"]), ("FEIN", c["fein"]), ("Entity type", c["entity"]), ("Mailing address", c["mailing"]),
          ("Years in business", str(c["years"])), ("NAICS", c["naics"]), ("Producer", c["broker"]), ("Producer contact", c["contact"])], cols=1, label_w=150)
    p.section("Description of operations")
    p.para(c["description"])
    tiv = sum(tiv_of(l) for l in c["locations"])
    eff = date.fromisoformat(c["effective"])
    p.section("Policy information")
    p.kv([("Proposed effective", c["effective"]), ("Proposed expiration", (eff.replace(year=eff.year + 1)).isoformat()), ("Coverage form", "Causes of Loss — Special (CP 10 30)"),
          ("Valuation", "Replacement cost"), ("Coinsurance", "Agreed value"), ("Limit requested", f"{money(tiv)} blanket"), ("Prior carrier", c["prior_carrier"]),
          ("Prior premium", money(c["prior_premium"]) if c["prior_premium"] else "Not disclosed")], cols=1, label_w=150)
    for i, l in enumerate(c["locations"], 1):
        stt = state_of(l)
        p.section(f"Premises {i} — {l['name']}")
        pairs = [("Premises address", loc_address(l, stt)), ("Occupancy", _v(l, "app", "occ_raw")), ("Construction", _v(l, "app", "cons_raw")),
                 ("Year built", str(_v(l, "app", "yb"))), ("Stories", str(_v(l, "app", "stories"))), ("Total area (sq ft)", f"{_v(l, 'app', 'sqft'):,}"),
                 ("Roof covering", l["roof_type"]), ("Roof year", str(_v(l, "app", "roof"))), ("Sprinklered", _spr_text(_v(l, "app", "spr"))),
                 ("Fire alarm", l["alarm"]), ("Protection class", str(l["ppc"])), ("Building limit", money(l["building"])), ("BPP limit", money(l["contents"] + l["stock"])),
                 ("Business income", money(l["bi"]))]
        st_ = _v(l, "app", "storage")
        if st_:
            pairs.append(("Max storage height (ft)", str(st_)))
        p.kv(pairs, cols=1, label_w=150)
    p.section("Loss history")
    shown = [x for x in c["losses"]]
    p.kv([("Losses in last 5 years", str(len(shown)) if c["loss_years"] >= 5 else f"{len(shown)} (last {c['loss_years']} years)"),
          ("Total incurred", money(sum(x["paid"] + x["reserve"] for x in shown)))], cols=1, label_w=150)
    p.section("Remarks")
    p.para("Applicant represents that the statements in this application are true and complete. Values are replacement cost as of the application date. "
           "Protection and construction as described by the insured.", size=8.5)
    p.signature(f"Authorized officer, {c['insured']}", "Applicant", signed)
    p.save()


# ============================================================================ SOV (same layouts the SOV parser reads)
def _spr_sov(v: float, tmpl: str):
    if tmpl == "B":
        return round(v * 100)
    if tmpl == "C":
        return "Full" if v >= 0.99 else "Partial" if v > 0 else "None"
    return "Y" if v >= 0.99 else "Partial" if v > 0 else "N"


def render_sov(c: dict, path: Path, as_of: str):
    tmpl = BROKER_TEMPLATE.get(c["broker"], "A")
    heads = list(TEMPLATES[tmpl])
    if any(l.get("storage") for l in c["locations"]):
        heads.insert(len(heads) - 1, "Max Storage Ht (ft)")
    wb = Workbook()
    ws = wb.active
    ws.title = "SOV" if tmpl != "C" else "Locations"
    ws["A1"] = f"{c['insured']} — Statement of Values"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Values as of {as_of} · new business submission"
    ws["A2"].font = Font(italic=True, size=10, color="555555")
    ws["A3"] = f"Prepared by {c['broker']}"
    ws["A3"].font = Font(size=9, color="777777")
    hr = 5
    for ci, h in enumerate(heads, 1):
        cell = ws.cell(row=hr, column=ci, value=h)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, Border(bottom=THIN)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width = max(9, min(28, len(h) + 4))
    ws.row_dimensions[hr].height = 30
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    for r, l in enumerate(c["locations"], hr + 1):
        stt = state_of(l)
        vals = {"Loc #": r - hr, "Loc No": r - hr, "Site": r - hr, "Bldg #": 1, "Building": 1,
                "Location Name": l["name"], "Site Name": l["name"], "Site Description": l["name"],
                "Street Address": l["address"], "Address": l["address"], "Location Address": l["address"],
                "City": l["city"], "State": stt, "ST": stt, "Zip": l["zip"], "ZIP": l["zip"], "Postal Code": l["zip"],
                "Occupancy": _v(l, "sov", "occ_raw"), "Occupancy Description": _v(l, "sov", "occ_raw"), "Use": _v(l, "sov", "occ_raw"),
                "Construction": _v(l, "sov", "cons_raw"), "Const Type": _v(l, "sov", "cons_raw"), "ISO Class / Construction": _v(l, "sov", "cons_raw"),
                "Year Built": _v(l, "sov", "yb"), "Yr Blt": _v(l, "sov", "yb"), "Year of Construction": _v(l, "sov", "yb"),
                "# Stories": l["stories"], "Stories": l["stories"], "No. of Floors": l["stories"],
                "Sq Ft": _v(l, "sov", "sqft"), "Area (sf)": _v(l, "sov", "sqft"), "Gross Area": _v(l, "sov", "sqft"),
                "Roof Year": _v(l, "sov", "roof"), "Roof Updated": _v(l, "sov", "roof"), "Roof Replaced": _v(l, "sov", "roof"),
                "Sprinklered": _spr_sov(_v(l, "sov", "spr"), tmpl), "Sprinkler %": _spr_sov(_v(l, "sov", "spr"), tmpl), "Sprinkler Protection": _spr_sov(_v(l, "sov", "spr"), tmpl),
                "Building Value": l["building"], "Bldg RC": l["building"], "Building (RC)": l["building"],
                "Contents / BPP": l["contents"], "BPP": l["contents"], "Contents": l["contents"],
                "Stock": l["stock"], "Inventory": l["stock"], "Stock / Inventory": l["stock"],
                "Business Income": l["bi"], "BI/EE (12 mo)": l["bi"], "Time Element": l["bi"],
                "TIV": tiv_of(l), "Total Insured Value": tiv_of(l), "Total": tiv_of(l), "Valuation": "RC",
                "Max Storage Ht (ft)": _v(l, "sov", "storage")}
        for ci, h in enumerate(heads, 1):
            v = vals.get(h)
            cell = ws.cell(row=r, column=ci, value=v)
            if isinstance(v, float):
                cell.number_format = "#,##0"
    nc = len(heads) + 1
    ws.cell(row=hr, column=nc, value="Internal notes").font = Font(bold=True, color="999999")
    ws.column_dimensions[get_column_letter(nc)].hidden = True
    s2 = wb.create_sheet("Broker notes")
    s2["A1"] = "Values supplied by insured; broker has not verified. Replacement cost basis."
    s2.sheet_state = "hidden"
    wb.save(path)


# ============================================================================ loss runs
def render_loss_run(c: dict, path: Path, valued: str, years: int, losses: list[dict], carrier: str):
    v = date.fromisoformat(valued)
    start = v.replace(year=v.year - years)
    p = Pdf(path, "carrier", "Loss Run — Commercial Property", carrier)
    p.title("Loss Run — Commercial Property", f"{carrier} · valued as of {valued}")
    p.kv([("Named insured", c["insured"]), ("Carrier", carrier), ("Experience period", f"{start.isoformat()} to {valued}"), ("Valuation date", valued),
          ("Policy years", str(years)), ("Claim count", str(len(losses))), ("Total incurred", money(sum(x["paid"] + x["reserve"] for x in losses)))], cols=1, label_w=150)
    if not losses:
        p.para("No losses reported during the experience period.")
    else:
        pre = "".join(w[0] for w in carrier.split()[:2]).upper()
        rows = []
        for i, x in enumerate(sorted(losses, key=lambda x: x["dol"])):
            l = c["locations"][x["loc"]]
            rows.append((f"{pre}-{x['dol'][2:4]}-{int(h(c['insured'], x['dol']), 16) % 9000 + 1000}", x["dol"], l["address"], x["cause"], x["status"], money(x["paid"]), money(x["reserve"]), money(x["paid"] + x["reserve"])))
        p.table(["Claim #", "Date of loss", "Location", "Cause", "Status", "Paid", "Reserve", "Incurred"], rows, [70, 60, 120, 70, 45, 55, 55, 60], align_right={5, 6, 7})
        p.section("Loss descriptions")
        for r_, x in zip(rows, sorted(losses, key=lambda x: x["dol"])):
            p.para(f"{r_[0]} — {x['desc']}", size=8.2)
    p.save()


def h(*parts) -> str:
    return hashlib.md5("|".join(map(str, parts)).encode()).hexdigest()[:6]


# ============================================================================ loss-control inspection
def render_inspection(c: dict, path: Path, insp: dict, verify: dict | None = None):
    title = "Loss Control Inspection Report" if not verify else "Verification Survey Report"
    p = Pdf(path, "engineering", title, f"{insp['firm'].split(' (')[0]} · {insp['date']}")
    p.title(title, f"Survey date {insp['date']} · {insp['engineer']} · {insp['firm']}")
    p.kv([("Insured", c["insured"]), ("Survey date", insp["date"]), ("Engineer", insp["engineer"]),
          ("Scope", "COPE, fire protection, natural hazard" if not verify else verify["scope"])], cols=1, label_w=150)
    for i, l in enumerate(c["locations"]):
        if verify and i not in verify["locs"]:
            continue
        stt = state_of(l)
        p.section(f"Location {i + 1} — {l['name']}")
        pairs = [("Address", loc_address(l, stt)), ("Construction", CONSTRUCTION[l["cons"]][0]), ("Year built", str(l["yb"])), ("Stories", str(l["stories"])),
                 ("Floor area (sq ft)", f"{l['sqft']:,}"), ("Roof covering", l["roof_type"]), ("Roof year (verified)", str(l["roof"])),
                 ("Occupancy observed", OCCUPANCY[l["occ"]][0]), ("Sprinkler protection", f"{l['spr'] * 100:.0f}% of area"), ("Fire alarm", l["alarm"]),
                 ("Public protection class", str(l["ppc"]))]
        design = l.get("design")
        if verify and verify.get("in_rack") and i == 0:
            design = l.get("storage")
        if l.get("storage"):
            pairs += [("Max storage height (ft)", str(l["storage"])), ("Sprinkler design storage height (ft)", str(design)), ("Commodity", l["commodity"])]
        if verify and verify.get("roof_condition") and i in verify.get("roof_locs", []):
            pairs.append(("Roof condition", verify["roof_condition"]))
        p.kv(pairs, cols=1, label_w=190)
        note = (verify or {}).get("notes", {}).get(i) or insp.get("notes", {}).get(i) or _narrative(l)
        p.para(note, size=8.5)
    p.signature(insp["engineer"], insp["firm"], insp["date"])
    p.save()


def _narrative(l: dict) -> str:
    s = f"The {l['name']} is a {l['stories']}-story {CONSTRUCTION[l['cons']][0].split(' (')[0].lower()} building of about {l['sqft']:,} sq ft built in {l['yb']}. "
    s += "Automatic sprinklers protect the full floor area. " if l["spr"] >= 0.99 else (f"Sprinklers protect about {l['spr'] * 100:.0f}% of the floor area. " if l["spr"] > 0 else "The building is not sprinklered. ")
    return s + "Housekeeping was good at the time of the visit."


# ============================================================================ manuscript wording
def render_manuscript(c: dict, path: Path, ms: dict, dated: str):
    p = Pdf(path, "broker", ms["title"], ms["form"])
    p.title(ms["title"], f"Form {ms['form']} · proposed by {c['broker']} · {dated}")
    p.kv([("Form number", ms["form"]), ("Named insured", c["insured"]), ("Attaches to", "CP 10 30 Causes of Loss — Special Form"), ("Edition", dated[:7])], cols=1, label_w=150)
    p.section("Wording")
    p.para(f"This endorsement modifies insurance provided under the Causes of Loss — Special Form. It is agreed that paragraph {ms['clause']}. "
           f"{ms['subject']} of the Causes of Loss — Special Form CP 10 30 is deleted in its entirety. All other terms and conditions remain unchanged.")
    p.save()


# ============================================================================ broker email
def render_email(c: dict, path: Path, sent: str, subject: str, body: str, attachments: list[tuple[str, Path, str]]):
    m = EmailMessage()
    m["From"] = f"{c['contact']} <{_email_of(c['contact'], c['broker'])}>"
    uw = USER_BY_ID[c["uw"]]["name"]
    m["To"] = f"{uw} <newbusiness@northgate.example>"
    m["Subject"] = subject
    m["Date"] = format_datetime(datetime.fromisoformat(sent + "T09:12:00-04:00"))
    m["Message-ID"] = f"<{hashlib.md5((subject + sent).encode()).hexdigest()[:16]}@broker.example>"
    m.set_content(body)
    for fname, p_, ctype in attachments:
        maintype, subtype = ctype.split("/")
        m.add_attachment(p_.read_bytes(), maintype=maintype, subtype=subtype, filename=fname)
    path.write_bytes(bytes(m))


def submission_body(c: dict) -> str:
    tiv = sum(tiv_of(l) for l in c["locations"])
    lines = [f"Hi {USER_BY_ID[c['uw']]['name'].split()[0]},", "",
             f"Please find attached our new business submission for {c['insured']}, with a requested effective date of {c['effective']}. "
             f"The schedule totals {money(tiv)} across {len(c['locations'])} location(s).", ""]
    if c.get("email"):
        lines += [c["email"], ""]
    if c.get("target"):
        lines += [f"The client has indicated a target premium of {money(c['target'])}.", ""]
    q = (date.fromisoformat(c["arrive"]) + timedelta(days=14)).isoformat()
    lines += [f"We would appreciate terms by {q}.", "", "Attached: ACORD application, statement of values, loss runs" +
              (", loss control inspection" if c.get("inspection") else "") + (", manuscript wording" if c.get("manuscript") else "") + ".", "",
              "Kind regards,", c["contact"], c["broker"]]
    return "\n".join(lines)


# ============================================================================ reference documents (new-business standards, authority matrix)
def render_standards(path: Path):
    p = Pdf(path, "carrier", "Northgate New Business Underwriting Standards", NB_STANDARDS["version"])
    p.title("New Business Underwriting Standards — Commercial Property", f"Version {NB_STANDARDS['version']} · effective {NB_STANDARDS['effective_from']} · supplements Guidelines 2026")
    p.para("These standards apply to new-business commercial property submissions and supplement the Northgate Commercial Property Underwriting Guidelines 2026. "
           "Each paragraph is implemented as a versioned rule in the Northgate rule library (product: decision).")
    for code, (title, text) in NB_SECTIONS.items():
        p.new_page()
        p.section(f"{code} {title}")
        p.para(text)
    p.save()


def render_authority(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "New business authority"
    heads = ["Level", "Role", "Pricing deviation (±)", "Max account TIV", "Max location TIV", "Max premium", "Guideline exceptions", "Declined-class exceptions"]
    for i, h_ in enumerate(heads, 1):
        cell = ws.cell(row=1, column=i, value=h_)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        ws.column_dimensions[get_column_letter(i)].width = 20
    for lvl, a in NB_AUTHORITY.items():
        ws.append([lvl, a["label"], a["max_price_dev"], a["max_account_tiv"], a["max_loc_tiv"], a["max_premium"], "Yes" if lvl >= 3 else "No", "Yes" if lvl == 4 else "No"])
        ws.cell(row=ws.max_row, column=3).number_format = "0%"
        for col in (4, 5, 6):
            ws.cell(row=ws.max_row, column=col).number_format = "$#,##0"
    wb.save(path)


def vendor_payload(c: dict, path: Path, enrich: list[dict], company: dict):
    path.write_text(json.dumps({"request": {"insured": c["insured"], "services": ["geocode", "hazard", "property", "valuation", "company"]},
                                "company": company, "locations": enrich}, indent=2))


# ============================================================================ broker replies and PAS documents
def render_roof_schedule(c: dict, path: Path, dated: str):
    p = Pdf(path, "broker", "Roof Replacement Schedule", c["insured"])
    p.title("Roof Replacement Schedule", f"Prepared by the insured's facilities team · {dated}")
    rows = []
    for l in c["locations"]:
        age = int(dated[:4]) - l["roof"]
        rows.append((l["name"], l["address"], str(l["roof"]), "Q2 2027 — full tear-off" if age > 20 else "No work planned"))
    p.table(["Location", "Address", "Roof year", "Planned replacement"], rows, [150, 150, 70, 130])
    p.para("Contract for the Fort Worth replacement signed with the roofing contractor; interim repairs completed after the 2022 hail event." if c["scenario"] == "D1" else
           "Schedule approved by the insured's facilities committee.", size=8.5)
    p.save()


def render_contractor_letter(c: dict, path: Path, dated: str):
    p = Pdf(path, "broker", "Sprinkler Contractor Completion Letter", c["insured"])
    l = c["locations"][0]
    p.title("In-Rack Sprinkler Installation — Completion Letter", f"Issued {dated} · Lowcountry Fire Protection (fictional)")
    p.kv([("Premises", loc_address(l, state_of(l))), ("Work", "In-rack sprinklers, two levels, racks 1–48"), ("Design standard", "NFPA 13 (2022), Group A plastics"),
          ("Completed", dated), ("Acceptance test", "Hydrostatic 200 psi, 2 hours — passed")], cols=1, label_w=150)
    p.save()


def render_pas_doc(kind: str, path: Path, c: dict, ref: str, dated: str, premium: float, act: dict, subjectivities: list[dict], prepared_by: str):
    from uwc import runtime_docs as RD
    eff = date.fromisoformat(c["effective"])
    term = (c["effective"], eff.replace(year=eff.year + 1).isoformat())
    terms = {"limit": act["limit"], "limit_basis": "blanket", "aop_deductible": act["aop"], "named_storm_ded_pct": act.get("ns_pct"),
             "named_storm_ded_min": act.get("ns_min"), "wind_hail_ded_pct": act.get("wh_pct"), "bi_sublimit": None, "flood_sublimit": None, "eq_sublimit": None}
    forms = ["CP 00 10", "CP 00 30", "CP 00 90", "CP 10 30", "IL 00 17"] + (["NS-CP 03 40"] if act.get("ns_pct") else []) + (["MS-WTR-01A"] if act.get("manuscript") == "limited" else ["MS-WTR-01"] if act.get("manuscript") else [])
    RD.render_contract(kind, path, name=c["insured"], broker=c["broker"], ref=ref, date=dated, term=term, premium=premium, terms=terms, forms=forms, safeguards=[],
                       subjectivities=subjectivities, version=1, prepared_by=prepared_by, share=act.get("line", 1.0),
                       layer="Primary" if act.get("line", 1.0) >= 1 else f"{act['line'] * 100:.0f}% part of {money(act['limit'] / act['line'])}")
    return CARRIER
