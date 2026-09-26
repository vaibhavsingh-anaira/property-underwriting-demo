"""World generator for the delegated book: what each coverholder actually wrote (the hidden plan), and the
files it sends — bordereaux in each coverholder's own layout, exposure returns, correction bordereaux.

The plan is generator truth. The engine never reads it: it only sees the xlsx files rendered from it,
the parsed BAA, the referral system and the claims bordereaux. Exceptions are injected at controlled
rates per coverholder story; everything else is written within authority."""
from __future__ import annotations

import hashlib
import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .refdata import (CLASSES, CONSTRUCTION, COVERHOLDERS, LAYOUTS, MESSY_TXN, MONTHS, NAME_A, NAME_B, OCC_TEXT, OUT_OF_SCHEDULE,
                      PREMIUM_LAYOUTS, RISK_FIELDS, PREMIUM_FIELDS, CLAIM_FIELDS, STREETS, SUFFIX, TXN_LABEL)

TAX = {"FL": 0.0494, "TX": 0.0485, "LA": 0.0485, "MS": 0.04, "AL": 0.06, "CA": 0.032, "NV": 0.035, "AZ": 0.03, "OR": 0.023, "PA": 0.03,
       "NJ": 0.05, "NY": 0.036, "MD": 0.03, "DE": 0.02, "OH": 0.05, "IN": 0.025, "MI": 0.02, "IL": 0.035, "WI": 0.03}

INJECT = {
    "ch_meridian": {"2026-05": [("limit", "correct"), ("pricing", "provide_ref")],
                    "2026-06": [("deductible", "endorse"), ("missing_ref_tier1", "dispute"), ("commission", "accept")],
                    "2026-07": [("showcase", "dispute"), ("deductible", "endorse"), ("class_bar", "cancel"), ("commission", "accept"), ("different_terms", "correct")],
                    "2026-08": [("limit", "dispute"), ("deductible", "endorse"), ("missing_ref_year", "provide_ref"), ("commission", "accept"), ("ns_ded", "correct"),
                                ("canx_return", "correct")]},
    "ch_northfield": {"2026-06": [("missing_ref_year", "provide_ref")], "2026-08": [("limit_stretch", "dispute")]},
    "ch_palmcoast": {"2026-06": [("ns_ded", "correct")], "2026-07": [("ns_ded", "endorse")]},
    "ch_ridgeway": {"2026-05": [("limit", "correct")], "2026-07": [("hidden_year", "dispute"), ("hidden_year", "provide_ref")]},
    "ch_sierra": {"2026-05": [("excluded_county", "admit")], "2026-07": [("pricing", "provide_ref")], "2026-08": [("prohibited_can", "cancel")]},
}


def _rng(*parts) -> random.Random:
    return random.Random(int(hashlib.md5("|".join(map(str, parts)).encode()).hexdigest()[:10], 16))


def _d(s: str) -> date:
    return date.fromisoformat(s)


def month_end(m: str) -> str:
    y, mo = map(int, m.split("-"))
    nxt = date(y + (mo == 12), mo % 12 + 1, 1)
    return (nxt - timedelta(days=1)).isoformat()


def _pick(rng, weights: dict):
    tot = sum(weights.values())
    x = rng.random() * tot
    for k, w in weights.items():
        x -= w
        if x <= 0:
            return k
    return k


def _terr(cfg, rng, allow_excluded=False, tier=None, zone=None):
    rows = [t for t in cfg["territories"] if (t[6] != "Excluded" or allow_excluded) and (tier is None or t[4] == tier) and (zone is None or t[5] == zone)]
    w = {i: max(0.1, t[7]) for i, t in enumerate(rows)}
    return rows[_pick(rng, w)]


def _name(rng, cls):
    return f"{rng.choice(NAME_A)} {rng.choice(NAME_B.get(cls, ['Holdings']))} {rng.choice(SUFFIX)}"


def _round_tiv(x):
    return float(round(x, -4))


def _ded_choice(rng, cfg, tiv, written):
    req = required_min_aop(cfg, tiv, written)
    opts = [d for d in (5_000, 10_000, 25_000, 50_000) if d >= req] or [req]
    return float(opts[0] if rng.random() < 0.7 else rng.choice(opts))


def required_min_aop(cfg, tiv, written):
    bands = list(cfg["min_aop"])
    e = cfg.get("endorsement", {})
    if e.get("min_aop") and written >= e["effective"]:
        cur = {a: (a, b, m) for a, b, m in bands}
        for a, b, m in e["min_aop"]:
            cur[a] = (a, cur.get(a, (a, b, m))[1], m)
        bands = sorted(cur.values())
    for a, b, m in bands:
        if tiv >= a and (b is None or tiv < b):
            return m
    return bands[-1][2]


def ns_min(cfg, tier, written):
    e = cfg.get("endorsement", {})
    base = dict(cfg["ns_min"])
    if e.get("ns_min") and written >= e["effective"]:
        base.update(e["ns_min"])
    return base.get(tier)


def commission_at(cfg, written):
    e = cfg.get("endorsement", {})
    return e["commission"] if e.get("commission") and written >= e["effective"] else cfg["commission"]


# ============================================================================ plan
def make_plan(ch: str) -> dict:
    cfg = COVERHOLDERS[ch]
    plan = {"rows": {}, "claims": [], "truth": {}, "exposure": [], "opening_gwp": 0.0, "pool": []}
    seq = [1]
    rng0 = _rng(ch, "pool")
    # policies bound Jan–Apr 2026 (not on the bordereaux we hold) that later endorse or cancel
    for i in range(60):
        plan["pool"].append(_risk(cfg, rng0, "2026-0%d" % rng0.randint(1, 4), "NEW", seq, pool=True))
    for m in MONTHS:
        rng = _rng(ch, m)
        rows = []
        for txn, n in cfg["rows"].items():
            for _ in range(n):
                if txn in ("NEW", "RENEWAL"):
                    rows.append(_risk(cfg, rng, m, txn, seq))
                else:
                    base = rng.choice(plan["pool"] + [r for mm in plan["rows"].values() for r in mm if r["transaction_type"] in ("NEW", "RENEWAL") and not r["_inj"]][-300:])
                    rows.append(_mid_term(cfg, rng, m, txn, base))
        if ch == "ch_northfield" and m == "2026-08":
            for k in range(4):  # larger risks after the requested capacity increase (referred if the amendment is not in force)
                r = _risk(cfg, rng, m, "NEW", seq)
                t = _round_tiv(rng.uniform(5_400_000, 7_200_000))
                r.update(tiv=t, limit=t, written_date=f"2026-08-{17 + k * 3:02d}", inception=f"2026-08-{20 + k * 3:02d}", expiry=f"2027-08-{20 + k * 3:02d}",
                         aop_deductible=25_000.0, _large=True)
                rows.append(r)
        for kind, resp in INJECT.get(ch, {}).get(m, []):
            _inject(cfg, rng, rows, kind, resp, plan["truth"], m, seq, plan["pool"])
        if cfg["layout"] == "messy":
            for r in rng.sample(rows, 3):
                r["_blank_state"] = True
            for r in rng.sample([x for x in rows if not x["_inj"]], 2):
                r["_tbd_tiv"] = True
            for r in rows:
                r["_bad_comm"] = rng.random() < 0.3
            for r in rng.sample([x for x in rows if not x["_inj"]], 2):
                r["_dup"] = True
        rows.sort(key=lambda r: (r["written_date"], r["certificate_ref"]))
        plan["rows"][m] = rows
    plan["claims"] = _claims(ch, cfg, plan)
    return plan


def _risk(cfg, rng, m, txn, seq, pool=False) -> dict:
    cls = _pick(rng, cfg["class_mix"])
    t = _terr(cfg, rng)
    if cfg.get("endorsement", {}).get("territories_add") and t[0] in {x[0] for x in cfg["endorsement"]["territories_add"]} and m < cfg["endorsement"]["effective"][:7]:
        t = _terr(cfg, rng)
    y, mo = map(int, m.split("-"))
    day = rng.randint(1, 28)
    written = date(y, mo, day)
    if txn == "RENEWAL":
        inception = written + timedelta(days=rng.randint(3, 20))
    else:
        inception = written + timedelta(days=rng.randint(0, 14))
    lo, hi = cfg["tiv_range"]
    tiv = _round_tiv(min(hi, max(lo, rng.lognormvariate(0, 0.55) * (lo + hi) / 3.2)))
    if t[4] == "Tier 1" and cfg["referral"].get("tiv_tier1") and tiv > cfg["referral"]["tiv_tier1"] and rng.random() < 0.8:
        tiv = _round_tiv(rng.uniform(lo, cfg["referral"]["tiv_tier1"] * 0.98))
    limit = tiv if rng.random() < 0.82 else _round_tiv(tiv * rng.uniform(0.5, 0.8))
    iso = _pick(rng, {1: 14, 2: 28, 3: 24, 4: 22, 6: 8} if t[0] != "FL" else {1: 4, 2: 16, 3: 18, 4: 44, 6: 18})
    yb = rng.randint(max(cfg["referral"]["year_built"] + 1, 1962), 2022)
    written_s = written.isoformat()
    nsm = ns_min(cfg, t[4], written_s)
    r = {"certificate_ref": f"{cfg['cert_prefix']}-26-{seq[0]:05d}", "transaction_type": txn, "insured_name": _name(rng, cls),
         "address": f"{rng.randint(100, 9899)} {rng.choice(STREETS)}", "city": t[2], "state": t[0], "county": t[1], "zip": f"{t[3]}{rng.randint(1, 99):02d}",
         "class_code": cls, "occupancy": rng.choice(OCC_TEXT[cls]), "construction": CONSTRUCTION[iso][0], "iso": iso, "year_built": yb,
         "tiv": tiv, "limit": limit, "aop_deductible": _ded_choice(rng, cfg, tiv, written_s), "ns_deductible_pct": (nsm if nsm and rng.random() < 0.8 else (nsm + 0.01 if nsm else None)),
         "written_date": written_s, "inception": inception.isoformat(), "expiry": (inception.replace(year=inception.year + 1)).isoformat(),
         "referral_ref": None, "prior_certificate": None, "prior_premium": None, "prior_tiv": None, "prior_aop_deductible": None,
         "_zone": t[5], "_tier": t[4], "_inj": None, "_rc": None}
    if pool:
        r["certificate_ref"] = f"{cfg['cert_prefix']}-26-{seq[0]:05d}"
    seq[0] += 1
    if txn == "RENEWAL":
        g = rng.gauss(0.055, 0.035)
        r["prior_tiv"] = _round_tiv(tiv / (1 + g))
        r["prior_certificate"] = f"{cfg['cert_prefix']}-25-{rng.randint(10000, 99999):05d}"
        r["prior_aop_deductible"] = r["aop_deductible"] if rng.random() < 0.85 else float(max(5_000, r["aop_deductible"] / 2))
        r["_rc"] = rng.gauss(cfg["rarc"].get(m, 0.0), 0.03)
    return r


def _mid_term(cfg, rng, m, txn, base) -> dict:
    y, mo = map(int, m.split("-"))
    eff = date(y, mo, rng.randint(1, 28))
    if eff.isoformat() < base["inception"]:
        eff = _d(base["inception"]) + timedelta(days=rng.randint(5, 30))
        if eff.isoformat()[:7] != m:
            eff = date(y, mo, 28)
    r = {k: v for k, v in base.items()}
    r.update(transaction_type=txn, written_date=eff.isoformat(), referral_ref=None, prior_certificate=None, prior_premium=None, prior_aop_deductible=None,
             _base_premium=base.get("_annual_premium"), _inj=None, _rc=None, _base=base["certificate_ref"])
    if txn == "ENDORSEMENT":
        r["prior_tiv"] = base["tiv"]
        new = _round_tiv(base["tiv"] * rng.uniform(1.04, 1.18) if rng.random() < 0.75 else base["tiv"] * rng.uniform(0.85, 0.96))
        r["tiv"] = new
        r["limit"] = min(new, cfg["max_limit"]) if base["limit"] >= base["tiv"] * 0.99 else base["limit"]
        if r["limit"] > cfg["max_limit"]:
            r["limit"] = cfg["max_limit"]
        r["aop_deductible"] = float(max(r["aop_deductible"], required_min_aop(cfg, r["tiv"], base["inception"])))
    else:
        r["prior_tiv"] = None
    return r


def _inject(cfg, rng, rows, kind, resp, truth, m, seq, pool):
    y, mo = map(int, m.split("-"))
    cands = [r for r in rows if r["transaction_type"] in ("NEW", "RENEWAL") and not r["_inj"] and not r.get("_large") and r["_tier"] != "Tier 1" and r["class_code"] != "HTL"]
    r = rng.choice(cands) if kind not in ("canx_return",) else None
    mx = cfg["max_limit"]
    info: dict = {"kind": kind, "response": resp, "month": m}
    if kind == "limit":
        r.update(tiv=_round_tiv(rng.uniform(mx * 1.15, mx * 1.45)))
        r["limit"] = r["tiv"]
        r["aop_deductible"] = float(required_min_aop(cfg, r["tiv"], r["written_date"]))
        info["true"] = {"limit": float(mx)}
    elif kind == "limit_stretch":
        r.update(tiv=8_200_000.0, limit=8_000_000.0, aop_deductible=25_000.0, written_date="2026-08-26", inception="2026-09-01", expiry="2027-09-01", class_code="WHS",
                 occupancy="Distribution warehouse", insured_name="Lakeshore Distribution Partners LLC")
    elif kind == "deductible":
        r["written_date"] = max(r["written_date"], f"{m}-02")
        r.update(tiv=_round_tiv(rng.uniform(5_300_000, 6_900_000)))
        r["limit"] = float(mx)
        req = required_min_aop(cfg, r["tiv"], r["written_date"])
        r["aop_deductible"] = float(req / 2)
        info["true"] = {"aop_deductible": float(req)}
    elif kind == "pricing":
        r["_adequacy"] = rng.uniform(0.80, 0.86)
        info["ref"] = True
    elif kind in ("missing_ref_tier1",):
        t = _terr(cfg, rng, tier="Tier 1")
        r.update(state=t[0], county=t[1], city=t[2], zip=f"{t[3]}{rng.randint(1, 99):02d}", _zone=t[5], _tier=t[4])
        r["tiv"] = _round_tiv(rng.uniform(cfg["referral"]["tiv_tier1"] * 1.25, cfg["referral"]["tiv_tier1"] * 1.8))
        r["limit"] = r["tiv"]
        r["aop_deductible"] = float(required_min_aop(cfg, r["tiv"], r["written_date"]))
        r["ns_deductible_pct"] = ns_min(cfg, "Tier 1", r["written_date"])
    elif kind in ("missing_ref_year", "hidden_year"):
        r["year_built"] = rng.randint(1938, cfg["referral"]["year_built"] - 3)
        info["ref"] = resp == "provide_ref"
    elif kind == "commission":
        info["true"] = {"commission_pct": commission_at(cfg, r["written_date"])}
        r["_commission"] = 0.25
    elif kind == "showcase":
        r.update(insured_name="Bayou City Cold Storage LLC", class_code="COL", occupancy="Refrigerated warehouse", construction="Joisted Masonry", iso=2,
                 state="TX", county="Harris", city="Houston", zip="77029", _zone="HOU", _tier="Tier 2", tiv=7_500_000.0, limit=7_500_000.0, aop_deductible=25_000.0,
                 ns_deductible_pct=0.02, year_built=1994, written_date="2026-07-14", inception="2026-07-15", expiry="2027-07-15", address="8120 Clinton Dr",
                 transaction_type="NEW", _adequacy=0.838, _commission=0.25)
    elif kind == "class_bar":
        r.update(class_code="BAR", occupancy="Bar & grill (late night)", insured_name=_name(rng, "BAR"))
    elif kind == "different_terms":
        r["tiv"] = _round_tiv(rng.uniform(cfg["referral"]["tiv_any"] * 1.12, cfg["referral"]["tiv_any"] * 1.25))
        r["limit"] = float(mx)
        r["aop_deductible"] = float(required_min_aop(cfg, r["tiv"], r["written_date"]))
        info["approved_tiv"] = _round_tiv(r["tiv"] * 0.9)
        info["true"] = {"tiv": info["approved_tiv"]}
        info["ref"] = True
    elif kind == "ns_ded":
        t = _terr(cfg, rng, tier="Tier 1")
        r.update(state=t[0], county=t[1], city=t[2], zip=f"{t[3]}{rng.randint(1, 99):02d}", _zone=t[5], _tier=t[4])
        if cfg["referral"].get("tiv_tier1") and r["tiv"] > cfg["referral"]["tiv_tier1"]:
            r["tiv"] = r["limit"] = _round_tiv(cfg["referral"]["tiv_tier1"] * 0.8)
        req = ns_min(cfg, "Tier 1", r["written_date"])
        r["ns_deductible_pct"] = round(req - (0.02 if req >= 0.05 else 0.01), 4)
        info["true"] = {"ns_deductible_pct": req}
    elif kind == "excluded_county":
        t = next(x for x in cfg["territories"] if x[1] == "Butte")
        r.update(insured_name="Paradise Ridge Supply Co.", class_code="RET", occupancy="Hardware store", state=t[0], county=t[1], city="Paradise", zip="95969",
                 _zone=t[5], _tier=t[4], tiv=2_140_000.0, limit=2_140_000.0, aop_deductible=10_000.0, construction="Frame", iso=1, year_built=1988,
                 written_date="2026-05-12", inception="2026-05-15", expiry="2027-05-15", transaction_type="NEW")
    elif kind == "prohibited_can":
        r.update(class_code="CAN", occupancy="Cannabis dispensary", insured_name="Green Canyon Wellness Dispensary LLC")
    elif kind == "canx_return":
        src = min((x for x in pool if x["inception"] < "2026-02-15"), key=lambda x: x["inception"])
        c = _mid_term(cfg, rng, m, "CANCELLATION", src)
        c["written_date"] = f"{m}-20"
        c["_inj"] = kind
        c["_full_return"] = True
        rows.append(c)
        truth[c["certificate_ref"] + "|CANCELLATION"] = {**info, "cert": c["certificate_ref"], "insured": c["insured_name"]}
        return
    r["_inj"] = kind
    truth[r["certificate_ref"]] = {**info, "cert": r["certificate_ref"], "insured": r["insured_name"]}


def price(rows: list[dict], cfg: dict, rated_fn, book: dict):
    """Premiums from the parsed authority's rating rules (rated_fn(row) -> technical premium). `book` maps
    certificate -> annual premium of policies already priced, so endorsements and cancellations are pro-rata of the real premium."""
    for r in rows:
        if r.get("_base") and book.get(r["_base"]):
            r["_base_premium"] = book[r["_base"]]
        rng = _rng(r["certificate_ref"], r["transaction_type"], r["written_date"])
        tech = rated_fn(r)
        if r["transaction_type"] in ("NEW", "RENEWAL"):
            adequacy = r.get("_adequacy") or rng.uniform(0.97, 1.12)
            r["gross_premium"] = float(round((tech or 5000) * adequacy))
            r["_annual_premium"] = r["gross_premium"]
            book[r["certificate_ref"]] = r["gross_premium"]
            if r["transaction_type"] == "RENEWAL":
                df = lambda d: next((f for dd, f in reversed([(5_000, 1.05), (10_000, 1.0), (25_000, 0.95), (50_000, 0.9), (100_000, 0.85)]) if d >= dd), 1.05)
                exp_f = r["tiv"] / r["prior_tiv"]
                terms_f = df(r["aop_deductible"]) / df(r["prior_aop_deductible"])
                r["prior_premium"] = float(round(r["gross_premium"] / (exp_f * terms_f * (1 + r["_rc"]))))
        elif r["transaction_type"] == "ENDORSEMENT":
            base_p = r.get("_base_premium") or (tech or 5000)
            frac = max(0.05, (_d(r["expiry"]) - _d(r["written_date"])).days / 365)
            r["gross_premium"] = float(round(base_p * (r["tiv"] / r["prior_tiv"] - 1) * frac))
        else:
            base_p = r.get("_base_premium") or (tech or 5000)
            frac = max(0.0, (_d(r["expiry"]) - _d(r["written_date"])).days / 365)
            r["gross_premium"] = float(-round(base_p * (1.0 if r.get("_full_return") else frac)))
            r["_base_premium"] = base_p
        pct = r.get("_commission") or commission_at(cfg, r["written_date"])
        r["commission_pct"] = pct
        r["commission_amount"] = round(r["gross_premium"] * pct, 2)
        r["taxes"] = round(r["gross_premium"] * TAX.get(r["state"], 0.03), 2)
        r["net_premium"] = round(r["gross_premium"] - r["commission_amount"], 2)


# ============================================================================ claims
def _claims(ch, cfg, plan) -> list[dict]:
    rng = _rng(ch, "claims")
    pols = [r for m in ("2026-05", "2026-06", "2026-07") for r in plan["rows"][m] if r["transaction_type"] in ("NEW", "RENEWAL") and not r["_inj"]]
    causes = [("Water damage", "Pipe burst", 0.3), ("Wind / hail", "Hail damage to roof", 0.25), ("Fire", "Kitchen fire", 0.1), ("Theft", "Break-in and theft of stock", 0.15),
              ("Water damage", "Roof leak", 0.1), ("Vandalism", "Vandalism to storefront", 0.1)]
    out = []
    n = {"ch_meridian": 14, "ch_northfield": 9, "ch_palmcoast": 7, "ch_ridgeway": 6, "ch_sierra": 10}[ch]
    for i in range(n):
        p = rng.choice(pols)
        cause, desc, _w = rng.choice(causes)
        dol = _d(p["inception"]) + timedelta(days=rng.randint(5, 60))
        if dol.isoformat() > "2026-07-28":
            dol = date(2026, 7, rng.randint(1, 28))
        if dol.isoformat() < p["inception"]:
            continue
        inc = round(rng.lognormvariate(9.6, 0.8), -2)
        rep = dol + timedelta(days=rng.randint(0, 6))
        out.append(_claim(cfg, len(out) + 1, p, dol.isoformat(), rep.isoformat(), cause, desc, [(rep.isoformat()[:7], 0.0, inc, "Open")], rng,
                          handler="Coverholder" if inc < 35_000 else "TPA — Westlake Claims"))
    if ch == "ch_sierra":
        pr = next(r for m in ("2026-05",) for r in plan["rows"][m] if r.get("_inj") == "excluded_county")
        out.append(_claim(cfg, 91, pr, "2026-07-18", "2026-07-19", "Wildfire", "Structure and stock destroyed in the Skyway Fire; total loss of building",
                          [("2026-07", 150_000.0, 1_200_000.0, "Open"), ("2026-08", 420_000.0, 1_480_000.0, "Open")], rng, cat="Skyway Fire (CA)"))
        rs = next(r for r in plan["rows"]["2026-05"] + plan["rows"]["2026-06"] if r["class_code"] == "RST" and r["state"] == "CA" and r["transaction_type"] in ("NEW", "RENEWAL") and not r["_inj"])
        rs["insured_name"] = "Casa Robles Kitchen LLC"
        out.append(_claim(cfg, 92, rs, "2026-06-09", "2026-06-10", "Fire", "Grease fire spread from cooking line to roof void; dining room closed",
                          [("2026-07", 85_000.0, 555_000.0, "Open"), ("2026-08", 240_000.0, 460_000.0, "Open")], rng, first_bdx="2026-07"))
        wh = next(r for r in plan["rows"]["2026-05"] if r["state"] == "AZ" and r["transaction_type"] in ("NEW", "RENEWAL") and not r["_inj"])
        out.append(_claim(cfg, 93, wh, "2026-06-02", "2026-06-03", "Water damage", "Fire sprinkler line failure over racked stock",
                          [("2026-06", 20_000.0, 160_000.0, "Open"), ("2026-07", 45_000.0, 475_000.0, "Open"), ("2026-08", 180_000.0, 350_000.0, "Open")], rng))
        sm = next(r for r in plan["rows"]["2026-06"] if r["transaction_type"] in ("NEW", "RENEWAL") and not r["_inj"] and r["state"] == "NV")
        out.append(_claim(cfg, 94, sm, "2026-06-20", "2026-06-21", "Theft", "Burglary — copper wiring and HVAC units stripped",
                          [("2026-06", 0.0, 78_000.0, "Open"), ("2026-07", 72_500.0, 0.0, "Closed")], rng, handler="Coverholder"))
    if ch == "ch_meridian":
        hp = next(r for r in plan["rows"]["2026-06"] if r["transaction_type"] in ("NEW", "RENEWAL") and not r["_inj"] and r["_tier"] == "Tier 2")
        out.append(_claim(cfg, 95, hp, "2026-07-08", "2026-07-09", "Wind / hail", "Hurricane Beryl-style wind event: roof membrane and rooftop units",
                          [("2026-07", 30_000.0, 210_000.0, "Open"), ("2026-08", 160_000.0, 95_000.0, "Open")], rng, cat="TS Delia (TX)"))
    return out


def _claim(cfg, i, pol, dol, rep, cause, desc, snaps, rng, cat=None, first_bdx=None, handler="TPA — Westlake Claims"):
    months = {}
    for m, paid, res, stt in snaps:
        months[m] = (paid, res, stt)
    first = first_bdx or rep[:7]
    for m in ("2026-05", "2026-06", "2026-07", "2026-08"):
        if m < first:
            continue
        if m not in months:
            prev = [x for x in sorted(months) if x < m]
            if prev:
                pp, rr, ss = months[prev[-1]]
                if ss == "Closed":
                    months[m] = (pp, 0.0, "Closed")
                else:
                    step = round(rr * rng.uniform(0.2, 0.5), -2)
                    months[m] = (pp + step, max(0.0, rr - step * rng.uniform(0.8, 1.2)), "Open" if rr - step > 2000 else "Closed")
    return {"claim_ref": f"{cfg['cert_prefix']}-C-{2600 + i:04d}", "certificate_ref": pol["certificate_ref"], "insured_name": pol["insured_name"], "state": pol["state"],
            "date_of_loss": dol, "date_reported": rep, "cause": cause, "description": desc, "handled_by": handler,
            "months": months, "first_bdx": first, "cat": cat}


# ============================================================================ files
HDR_FILL = PatternFill("solid", fgColor="1F3A5F")
HDR_FONT = Font(bold=True, color="FFFFFF")


def _fmt_date(v, messy):
    if v is None:
        return None
    if messy:
        return datetime.fromisoformat(v).strftime("%m/%d/%y")
    return v[:10]


def write_risk(path: Path, cfg: dict, month: str, rows: list[dict], layout: str, kind: str = "risk", title: str | None = None):
    """Risk or premium bordereau in the coverholder's own template (or the CRS template)."""
    fields = RISK_FIELDS if kind == "risk" else PREMIUM_FIELDS
    lay = (LAYOUTS if kind == "risk" else PREMIUM_LAYOUTS).get(layout) or {f: fields[f][0] for f in fields}
    messy = layout == "messy"
    wb = Workbook()
    ws = wb.active
    ws.title = {"risk": "Risk Bordereau", "premium": "Premium Bordereau"}[kind] if not messy else ("Production Report" if kind == "risk" else "Premium Report")
    start = 1
    if messy:
        ws["A1"] = f"{cfg['name']} — Monthly {'Production' if kind == 'risk' else 'Premium'} Report"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = f"Program: {cfg['program']} · Month: {datetime.fromisoformat(month + '-01').strftime('%B %Y')}"
        ws["A3"] = "Prepared by Ridgeway Programs operations — internal format"
        start = 5
    elif title:
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=12)
        start = 3
    cols = [f for f in fields if f in lay]
    for ci, f in enumerate(cols, 1):
        c = ws.cell(row=start, column=ci, value=lay[f])
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = max(11, min(28, len(lay[f]) + 4))
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    pe = month_end(month)
    r_i = start + 1
    for r in rows:
        for ci, f in enumerate(cols, 1):
            v = _value(r, f, cfg, pe, messy)
            c = ws.cell(row=r_i, column=ci, value=v)
            typ = fields[f][1]
            if typ == "money" and isinstance(v, (int, float)):
                c.number_format = "#,##0"
            elif typ == "pct" and isinstance(v, float) and v < 1:
                c.number_format = "0.0%"
            elif typ == "date" and not isinstance(v, str) and v is not None:
                c.number_format = "yyyy-mm-dd"
        r_i += 1
        if r.get("_dup"):
            for ci, f in enumerate(cols, 1):
                ws.cell(row=r_i, column=ci, value=_value(r, f, cfg, pe, messy))
            r_i += 1
    if messy:
        ws.cell(row=r_i + 1, column=1, value="TOTAL")
        gcol = cols.index("gross_premium") + 1 if "gross_premium" in cols else None
        if gcol:
            ws.cell(row=r_i + 1, column=gcol, value=round(sum(x.get("gross_premium") or 0 for x in rows)))
        tcol = cols.index("tiv") + 1 if "tiv" in cols else None
        if tcol:
            ws.cell(row=r_i + 1, column=tcol, value=round(sum(x.get("tiv") or 0 for x in rows if isinstance(x.get("tiv"), (int, float)))))
    wb.save(path)


def _value(r, f, cfg, pe, messy):
    if f == "umr":
        return cfg["umr"]
    if f == "agreement_no":
        return cfg["agreement"]
    if f == "reporting_period":
        return _fmt_date(pe, messy) if not messy else datetime.fromisoformat(pe).strftime("%b-%y")
    v = r.get(f)
    if f == "transaction_type":
        return MESSY_TXN[v] if messy else TXN_LABEL[v]
    if messy:
        if f in ("written_date", "inception", "expiry"):
            return _fmt_date(v, True)
        if f == "commission_pct" and v is not None:
            return round(v * 100, 1)
        if f == "state" and r.get("_blank_state"):
            return None
        if f == "tiv" and r.get("_tbd_tiv"):
            return "TBD"
        if f == "commission_amount" and r.get("_bad_comm"):
            return round((r["gross_premium"] - r["taxes"]) * r["commission_pct"], 2)
        if f == "net_premium" and r.get("_bad_comm"):
            return round(r["gross_premium"] - (r["gross_premium"] - r["taxes"]) * r["commission_pct"], 2)
        if f == "zip" and v:
            return int(v)
    if f in ("written_date", "inception", "expiry", "date_of_loss", "date_reported"):
        return _fmt_date(v, False)
    if f == "ns_deductible_pct" and v is None:
        return None
    return v


def write_claims(path: Path, cfg: dict, month: str, claims: list[dict], layout: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Claims Bordereau"
    cols = list(CLAIM_FIELDS)
    for ci, f in enumerate(cols, 1):
        c = ws.cell(row=1, column=ci, value=CLAIM_FIELDS[f][0] if layout != "variant" else {"claim_ref": "Claim Number", "date_of_loss": "DOL", "reserve": "O/S Reserve"}.get(f, CLAIM_FIELDS[f][0]))
        c.font, c.fill = HDR_FONT, HDR_FILL
        ws.column_dimensions[c.column_letter].width = 16 if f != "description" else 44
    ws.freeze_panes = "A2"
    pe = month_end(month)
    ri = 2
    for cl in claims:
        if cl["first_bdx"] > month or month not in cl["months"]:
            continue
        paid, res, stt = cl["months"][month]
        vals = {"umr": cfg["umr"], "agreement_no": cfg["agreement"], "reporting_period": pe, "claim_ref": cl["claim_ref"], "certificate_ref": cl["certificate_ref"],
                "insured_name": cl["insured_name"], "state": cl["state"], "date_of_loss": cl["date_of_loss"], "date_reported": cl["date_reported"],
                "cause": cl["cause"], "description": cl["description"] + (f" · CAT: {cl['cat']}" if cl.get("cat") else ""), "status": stt, "paid": paid, "reserve": res,
                "incurred": paid + res, "handled_by": cl["handled_by"]}
        for ci, f in enumerate(cols, 1):
            c = ws.cell(row=ri, column=ci, value=vals[f])
            if CLAIM_FIELDS[f][1] == "money":
                c.number_format = "#,##0"
            elif CLAIM_FIELDS[f][1] == "date":
                c.number_format = "yyyy-mm-dd"
        ri += 1
    wb.save(path)


def write_exposure(path: Path, cfg: dict, rows: list[dict], gwp: float):
    wb = Workbook()
    ws = wb.active
    ws.title = "Exposure by County"
    ws["A1"] = f"{cfg['name']} — in-force exposure return as at 30 Apr 2026"
    ws["A1"].font = Font(bold=True, size=12)
    heads = ["Risk State", "Risk County", "Policies In Force", "In-force TIV"]
    for ci, h in enumerate(heads, 1):
        c = ws.cell(row=3, column=ci, value=h)
        c.font, c.fill = HDR_FONT, HDR_FILL
        ws.column_dimensions[c.column_letter].width = 20
    for i, r in enumerate(rows, 4):
        ws.cell(row=i, column=1, value=r["state"])
        ws.cell(row=i, column=2, value=r["county"])
        ws.cell(row=i, column=3, value=r["policies"])
        ws.cell(row=i, column=4, value=r["tiv"]).number_format = "#,##0"
    s = wb.create_sheet("Summary")
    s["A1"], s["B1"] = "Gross written premium 1 Jan – 30 Apr 2026", round(gwp)
    s["B1"].number_format = "#,##0"
    s["A2"], s["B2"] = "Policies in force", sum(r["policies"] for r in rows)
    wb.save(path)
